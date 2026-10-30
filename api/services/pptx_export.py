"""PowerPoint export of a process model (L1 end-to-end process down to L5 activity).

Mirrors the Signavio BPMN export (api/routes/process_export.py) but renders the same
``ProcessModelDocument`` into slides: a summary deck, one slide per L1 (its L2/L3/L4 tree)
and one slide per L4 (its L5 activities with rule and table coverage).
"""

from __future__ import annotations

import io
from typing import Optional

from pptx import Presentation
from pptx.util import Inches, Pt

from api.models.process_model import L1, L2, L3, L4, ProcessModelDocument

_DARK = "1F2937"
_MUTED = "6B7280"
_ACCENT = "0F766E"


def _add_title(prs: Presentation, text: str, sub: Optional[str] = None) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    box = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(12.1), Inches(0.9))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = text
    run.font.size = Pt(30)
    run.font.bold = True
    run.font.color.rgb = _hex(_DARK)
    if sub:
        sub_box = slide.shapes.add_textbox(Inches(0.6), Inches(1.2), Inches(12.1), Inches(0.5))
        srun = sub_box.text_frame.paragraphs[0].add_run()
        srun.text = sub
        srun.font.size = Pt(14)
        srun.font.color.rgb = _hex(_MUTED)


def _add_body(prs: Presentation, lines: list[tuple[str, int, str]], top: float = 1.9) -> None:
    """lines: (text, indent level 0..3, colour hex or "")."""
    box = prs.slides[-1].shapes.add_textbox(Inches(0.6), Inches(top), Inches(12.1), Inches(5.2))
    tf = box.text_frame
    tf.word_wrap = True
    for text, level, colour in lines:
        p = tf.add_paragraph() if tf.paragraphs[0].runs else tf.paragraphs[0]
        p.level = level
        run = p.add_run()
        run.text = text
        run.font.size = Pt(16 - 2 * level)
        run.font.color.rgb = _hex(colour or _DARK)


def _hex(s: str):
    from pptx.dml.color import RGBColor
    return RGBColor(int(s[0:2], 16), int(s[2:4], 16), int(s[4:6], 16))


def _title_slide(prs: Presentation, meta: dict) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(11.7), Inches(2.0))
    p = box.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = meta.get("name") or "Process model"
    run.font.size = Pt(40)
    run.font.bold = True
    run.font.color.rgb = _hex(_ACCENT)
    sub = box.text_frame.add_paragraph()
    srun = sub.add_run()
    srun.text = f"L1–L5 business process, version {meta.get('version_no', 0)}, exported {meta.get('exported_at', '')}"
    srun.font.size = Pt(14)
    srun.font.color.rgb = _hex(_MUTED)


def _l1_slide(prs: Presentation, l1: L1, doc: ProcessModelDocument) -> None:
    lines: list[tuple[str, int, str]] = []
    for l2 in l1.l2:
        lines.append((l2.name, 0, _DARK))
        for l3 in l2.l3:
            lines.append((f"{l3.name}, {len(l3.l4)} sub-processes", 1, _MUTED))
            for l4 in l3.l4:
                lines.append((f"{l4.name}, {len(l4.activities)} activities", 2, _ACCENT))
    n_l5 = sum(len(l4.activities) for a in l1.l2 for b in a.l3 for l4 in b.l4)
    _add_title(prs, l1.name, f"{len(l1.l2)} process groups, {n_l5} activities")
    _add_body(prs, lines or [("No process groups", 0, _MUTED)])


def _l4_slide(prs: Presentation, l1: L1, l2: L2, l3: L3, l4: L4) -> None:
    lines: list[tuple[str, int, str]] = []
    for a in l4.activities:
        rule = f"{len(a.check_ids)} rules, {len(a.sap_tables)} tables" if (a.check_ids or a.sap_tables) else "no DQ coverage"
        lines.append((f"{a.name}  —  {rule}", 0, _DARK))
        if a.tcode:
            lines.append((f"t-code {a.tcode}", 1, _MUTED))
    _add_title(prs, l4.name, f"{l1.name} → {l2.name} → {l3.name}" + (f", t-code {l4.tcode}" if l4.tcode else ""))
    _add_body(prs, lines or [("No activities", 0, _MUTED)])


def export_pptx(doc: ProcessModelDocument, meta: dict) -> bytes:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    _title_slide(prs, meta)
    for l1 in doc.l1:
        _l1_slide(prs, l1, doc)
        for l2 in l1.l2:
            for l3 in l2.l3:
                for l4 in l3.l4:
                    _l4_slide(prs, l1, l2, l3, l4)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
