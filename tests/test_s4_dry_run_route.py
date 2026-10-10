"""Worker mode s4_dry_run: load_sim findings are folded into the module verdict and
persisted alongside the ordinary transfer gaps.

Runs with MERIDIAN_TEST_DB_URL, same skip convention as test_insights_impact_route.py.
"""

from __future__ import annotations

import io
import json
import os
import uuid

import pandas as pd
import pytest
from httpx import ASGITransport, AsyncClient
from sqlalchemy import create_engine, text

from api.deps import Tenant, get_tenant
from api.main import app
from api import deps as api_deps

pytestmark = pytest.mark.skipif(
    not os.getenv("MERIDIAN_TEST_DB_URL"), reason="requires MERIDIAN_TEST_DB_URL"
)


@pytest.fixture
def seeded():
    engine = create_engine(os.environ["MERIDIAN_TEST_DB_URL"])
    tid, sid, rid, vid = (str(uuid.uuid4()) for _ in range(4))
    with engine.begin() as c:
        c.execute(text("INSERT INTO tenants (id, name) VALUES (:t, :t)"), {"t": tid})
        c.execute(text("INSERT INTO sap_systems (id, tenant_id, name, system_type) VALUES "
                       "(:s, :t, 'PRD', 'ecc6')"), {"s": sid, "t": tid})
        c.execute(text("INSERT INTO analysis_versions (id, tenant_id, status) VALUES (:v, :t, 'complete')"),
                  {"v": vid, "t": tid})
        c.execute(text("INSERT INTO migration_runs (id, tenant_id, mode, source_system_id, modules, status) "
                       "VALUES (:r, :t, 's4_dry_run', :s, '{material_master}', 'queued')"),
                  {"r": rid, "t": tid, "s": sid})
    yield {"tid": tid, "sid": sid, "rid": rid, "vid": vid, "engine": engine}
    with engine.begin() as c:
        c.execute(text("DELETE FROM migration_gap_findings WHERE tenant_id = :t"), {"t": tid})
        c.execute(text("DELETE FROM transfer_field_mappings WHERE tenant_id = :t"), {"t": tid})
        c.execute(text("DELETE FROM migration_runs WHERE tenant_id = :t"), {"t": tid})
        c.execute(text("DELETE FROM analysis_versions WHERE tenant_id = :t"), {"t": tid})
        c.execute(text("DELETE FROM sap_systems WHERE tenant_id = :t"), {"t": tid})
        c.execute(text("DELETE FROM tenants WHERE id = :t"), {"t": tid})
    engine.dispose()


def test_s4_dry_run_folds_load_sim_into_verdict(seeded, monkeypatch):
    from checks.frames import TableFrames
    from sap.ddic import get_dictionary
    from workers.tasks import run_migration as mod

    from api.services.migration import load_sim

    # Two MARA rows that collide after ALPHA conversion: "123" and "0000123" both
    # resolve to the same internal material key. Production shape: TABLE.FIELD columns,
    # so the engine keys records MATNR=.. exactly as the sim does and fold() can match them.
    mara = pd.DataFrame({"MARA.MATNR": ["123", "0000123", "OK-1"], "MARA.MTART": ["ROH"] * 3,
                         "MARA.MBRSH": ["M"] * 3, "MARA.MEINS": ["EA"] * 3})
    frames = TableFrames({"MARA": mara}, get_dictionary("s4hana"))

    monkeypatch.setattr(mod, "resolve_source_version",
                        lambda session, source_system_id, source_version_id: (seeded["vid"], {"dataset_path": "x/"}))
    monkeypatch.setattr("workers.dataset.load_dataset",
                        lambda *a, **k: (frames, None, len(mara), len(mara.columns)))

    # The test DB role bypasses RLS, so load_mappings would see other tenants' stored mappings;
    # seed fresh from this frame instead.
    monkeypatch.setattr(mod, "load_mappings", lambda session, module, target_type: [])

    # The engine alone: no load-sim findings. These records pass it.
    real_simulate = load_sim.simulate
    monkeypatch.setattr(load_sim, "simulate", lambda *a, **k: [])
    out = mod.run_migration.run(seeded["tid"], seeded["rid"], "s4_dry_run", None, None, ["material_master"])
    assert out["status"] == "analysed", out
    with seeded["engine"].begin() as c:
        engine_only = c.execute(text("SELECT gap_summary FROM migration_runs WHERE id = :r"),
                                {"r": seeded["rid"]}).scalar()["material_master"]
    assert engine_only["records"] == 3 and engine_only["blocked_records"] == 0, engine_only
    monkeypatch.setattr(load_sim, "simulate", real_simulate)

    out = mod.run_migration.run(seeded["tid"], seeded["rid"], "s4_dry_run", None, None,
                                ["material_master"])
    assert out["status"] == "analysed", out

    engine = seeded["engine"]
    with engine.begin() as c:
        rows = c.execute(text("SELECT gap_type, detail FROM migration_gap_findings WHERE run_id = :r "
                              "AND gap_type = 's4_load'"), {"r": seeded["rid"]}).fetchall()
        gap_summary = c.execute(text("SELECT gap_summary FROM migration_runs WHERE id = :r"),
                                {"r": seeded["rid"]}).scalar()

    assert len(rows) == 2
    assert all(d.startswith("S4L-MM-MATNR-ALPHA") for _, d in rows)
    # the fold: the two sim-only keys are now blocked, the clean OK-1 is not
    assert gap_summary["material_master"]["blocked_records"] == 2
    assert gap_summary["material_master"]["verdict"] == "no-go"
    assert gap_summary["material_master"]["mode"] == "s4_dry_run"
    assert gap_summary["material_master"]["s4_load"]["S4L-MM-MATNR-ALPHA"] == 2

    with engine.begin() as c:
        total_before = c.execute(text("SELECT count(*) FROM migration_gap_findings WHERE run_id = :r"),
                                 {"r": seeded["rid"]}).scalar()

    # Idempotency: retrying the same run_id must not duplicate gap rows.
    out2 = mod.run_migration.run(seeded["tid"], seeded["rid"], "s4_dry_run", None, None,
                                 ["material_master"])
    assert out2["status"] == "analysed", out2
    with engine.begin() as c:
        total_after = c.execute(text("SELECT count(*) FROM migration_gap_findings WHERE run_id = :r"),
                                {"r": seeded["rid"]}).scalar()
        rows_after = c.execute(text("SELECT gap_type, detail FROM migration_gap_findings WHERE run_id = :r "
                                    "AND gap_type = 's4_load'"), {"r": seeded["rid"]}).fetchall()
    assert total_after == total_before
    assert len(rows_after) == len(rows)


# ── Routes over the dry-run results (task 9) ─────────────────────────────────


@pytest.fixture
def analysed_dry_run():
    engine = create_engine(os.environ["MERIDIAN_TEST_DB_URL"])
    t1, t2, sid, rid, vid = (str(uuid.uuid4()) for _ in range(5))
    gap_summary = {
        "material_master": {
            "records": 2, "blocked_records": 1, "verdict": "no-go", "score": 50.0,
            "mode": "s4_dry_run", "s4_load": {"S4L-MM-MATNR-ALPHA": 2},
        }
    }
    with engine.begin() as c:
        for t in (t1, t2):
            c.execute(text("INSERT INTO tenants (id, name) VALUES (:t, :t)"), {"t": t})
        c.execute(text("INSERT INTO sap_systems (id, tenant_id, name, system_type) VALUES "
                       "(:s, :t, 'PRD', 'ecc6')"), {"s": sid, "t": t1})
        c.execute(text("INSERT INTO analysis_versions (id, tenant_id, status) VALUES (:v, :t, 'complete')"),
                  {"v": vid, "t": t1})
        c.execute(text("""
            INSERT INTO migration_runs (id, tenant_id, mode, source_system_id, modules, status,
                                        gap_summary, readiness_verdict, readiness_score, completed_at)
            VALUES (:r, :t, 's4_dry_run', :s, '{material_master}', 'analysed',
                    CAST(:gs AS jsonb), 'no-go', 50.0, now())
        """), {"r": rid, "t": t1, "s": sid, "gs": json.dumps(gap_summary)})
        c.execute(text("""
            INSERT INTO migration_gap_findings (tenant_id, run_id, module, object_type, record_key,
                                                 gap_type, severity, detail)
            VALUES
            (:t, :r, 'material_master', 'MARA', 'MATNR1', 's4_load', 'critical', 'S4L-MM-MATNR-ALPHA collision 1'),
            (:t, :r, 'material_master', 'MARA', 'MATNR1', 's4_load', 'high', 'S4L-MM-MATNR-ALPHA collision 2'),
            (:t, :r, 'material_master', 'MARA', 'MATNR2', 's4_load', 'medium', 'S4L-MM-OTHER minor note')
        """), {"t": t1, "r": rid})
    yield {"t1": t1, "t2": t2, "rid": rid, "sid": sid, "vid": vid}
    with engine.begin() as c:
        c.execute(text("DELETE FROM migration_gap_findings WHERE tenant_id = :t"), {"t": t1})
        c.execute(text("DELETE FROM migration_runs WHERE tenant_id = :t"), {"t": t1})
        c.execute(text("DELETE FROM analysis_versions WHERE tenant_id = :t"), {"t": t1})
        c.execute(text("DELETE FROM sap_systems WHERE tenant_id = :t"), {"t": t1})
        c.execute(text("DELETE FROM tenants WHERE id IN (:t1, :t2)"), {"t1": t1, "t2": t2})
    engine.dispose()


def _patch_tenant(monkeypatch, tenant_id: str):
    monkeypatch.setattr("api.middleware.local_auth._load_jwt_secret", lambda: "test-secret")
    monkeypatch.setattr(
        "api.middleware.local_auth.decode_access_token",
        lambda token, secret: {
            "sub": "00000000-0000-0000-0000-000000000002",
            "email": "dev@example.com",
            "role": "admin",
        },
    )
    monkeypatch.setenv("MERIDIAN_DEV_ROLE_HEADER", "1")
    monkeypatch.setitem(
        app.dependency_overrides, get_tenant,
        lambda: Tenant(uuid.UUID(tenant_id), "T", []),
    )


_HEADERS = {"X-User-Role": "admin", "Authorization": "Bearer test-token"}


@pytest.mark.anyio
async def test_records_endpoint_returns_only_load_fail(analysed_dry_run, monkeypatch):
    rid = analysed_dry_run["rid"]
    _patch_tenant(monkeypatch, analysed_dry_run["t1"])
    await api_deps.engine.dispose()
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        r = await client.get(f"/api/v1/migration/runs/{rid}/records?status=load_fail", headers=_HEADERS)
    assert r.status_code == 200
    body = r.json()
    assert body["total"] == 1
    assert len(body["rows"]) == 1
    row = body["rows"][0]
    assert row["record_key"] == "MATNR1"
    assert row["status"] == "load_fail"
    assert row["module"] == "material_master"
    assert all(reason.startswith("S4L-") for reason in row["reasons"])


@pytest.mark.anyio
async def test_records_endpoint_is_tenant_isolated(analysed_dry_run, monkeypatch):
    rid = analysed_dry_run["rid"]
    _patch_tenant(monkeypatch, analysed_dry_run["t2"])
    await api_deps.engine.dispose()
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        r = await client.get(f"/api/v1/migration/runs/{rid}/records", headers=_HEADERS)
    assert r.status_code == 404


@pytest.mark.anyio
async def test_dry_run_xlsx_export(analysed_dry_run, monkeypatch):
    import openpyxl

    rid = analysed_dry_run["rid"]
    _patch_tenant(monkeypatch, analysed_dry_run["t1"])
    await api_deps.engine.dispose()
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        r = await client.get(f"/api/v1/migration/runs/{rid}/dry-run/xlsx", headers=_HEADERS)
    assert r.status_code == 200
    assert r.content[:2] == b"PK"
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert "Load fail" in wb.sheetnames


@pytest.mark.anyio
async def test_dry_run_xlsx_export_escapes_formula_injection(analysed_dry_run, monkeypatch):
    """A gap detail/record_key that starts with a formula char must come out of the xlsx
    export prefixed with a single quote (Excel/Sheets formula-injection guard)."""
    import openpyxl

    rid = analysed_dry_run["rid"]
    t1 = analysed_dry_run["t1"]
    engine = create_engine(os.environ["MERIDIAN_TEST_DB_URL"])
    with engine.begin() as c:
        c.execute(text("""
            INSERT INTO migration_gap_findings (tenant_id, run_id, module, object_type, record_key,
                                                 gap_type, severity, detail)
            VALUES (:t, :r, 'material_master', 'MARA', '=cmd|/C calc', 's4_load', 'critical',
                    '=HYPERLINK("http://evil") formula payload')
        """), {"t": t1, "r": rid})
    engine.dispose()

    _patch_tenant(monkeypatch, t1)
    await api_deps.engine.dispose()
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        r = await client.get(f"/api/v1/migration/runs/{rid}/dry-run/xlsx", headers=_HEADERS)
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    ws = wb["Load fail"]
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    rk_col = header.index("record_key")
    reason_col = header.index("reasons")
    rows = [[c.value for c in row] for row in ws.iter_rows(min_row=2)]
    hit = next(row for row in rows if "cmd" in str(row[rk_col]))
    assert hit[rk_col].startswith("'=")
    assert hit[reason_col].startswith("'=")


def test_to_xlsx_sanitize_formulas_flag():
    """Default to_xlsx leaves every value byte-for-byte (the /export SAP load file route relies
    on this for real '+' phone numbers and '-' balances). sanitize_formulas=True escapes formula
    every formula prefix through the shared export guard."""
    import pandas as pd
    import openpyxl

    from api.services.migration.export import to_xlsx

    df = pd.DataFrame({"v": ["=1+1", "+27115551234", "-5", "-1.2", "@mention", "plain"]})

    default_wb = openpyxl.load_workbook(io.BytesIO(to_xlsx({"Sheet": df})))
    default_vals = [c[0].value for c in default_wb["Sheet"].iter_rows(min_row=2)]
    assert default_vals == ["=1+1", "+27115551234", "-5", "-1.2", "@mention", "plain"]

    sanitized_wb = openpyxl.load_workbook(io.BytesIO(to_xlsx({"Sheet": df}, sanitize_formulas=True)))
    sanitized_vals = [c[0].value for c in sanitized_wb["Sheet"].iter_rows(min_row=2)]
    # Shared guard_frame: a human report, so every prefix is escaped (numeric columns stay numbers).
    assert sanitized_vals == ["'=1+1", "'+27115551234", "'-5", "'-1.2", "'@mention", "plain"]


@pytest.mark.anyio
async def test_dry_run_pdf_export(analysed_dry_run, monkeypatch):
    pytest.importorskip("weasyprint")

    rid = analysed_dry_run["rid"]
    _patch_tenant(monkeypatch, analysed_dry_run["t1"])
    await api_deps.engine.dispose()
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        r = await client.get(f"/api/v1/migration/runs/{rid}/dry-run/pdf", headers=_HEADERS)
    assert r.status_code == 200
    assert r.headers["content-type"] == "application/pdf"


@pytest.mark.anyio
async def test_analyze_s4_dry_run_rejects_dest_system_id(analysed_dry_run, monkeypatch):
    _patch_tenant(monkeypatch, analysed_dry_run["t1"])
    await api_deps.engine.dispose()
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        r = await client.post(
            "/api/v1/migration/analyze",
            json={"mode": "s4_dry_run", "source_system_id": None, "source_version_id": str(uuid.uuid4()),
                  "dest_system_id": str(uuid.uuid4()), "modules": ["material_master"]},
            headers=_HEADERS,
        )
    assert r.status_code == 400


@pytest.mark.anyio
async def test_analyze_s4_dry_run_queues_run_migration(analysed_dry_run, monkeypatch):
    from workers.tasks import run_migration as mod

    calls: list[tuple[object, ...]] = []

    class _Task:
        id = "task-1"

    def _delay(*args: object) -> _Task:
        calls.append(args)
        return _Task()

    monkeypatch.setattr(mod.run_migration, "delay", _delay)
    monkeypatch.setattr("api.routes.migration.current_user_id", lambda request: None)
    _patch_tenant(monkeypatch, analysed_dry_run["t1"])
    await api_deps.engine.dispose()
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        r = await client.post(
            "/api/v1/migration/analyze",
            json={"mode": "s4_dry_run", "source_system_id": analysed_dry_run["sid"],
                  "source_version_id": analysed_dry_run["vid"], "modules": ["material_master"]},
            headers=_HEADERS,
        )
    assert r.status_code in (200, 202), r.text
    body = r.json()
    assert body["status"] == "queued" and body["mode"] == "s4_dry_run" and body["task_id"] == "task-1"
    assert calls == [(analysed_dry_run["t1"], body["run_id"], "s4_dry_run", analysed_dry_run["sid"], None,
                      ["material_master"], analysed_dry_run["vid"], "s4hana")]
    engine = create_engine(os.environ["MERIDIAN_TEST_DB_URL"])
    with engine.begin() as c:
        row = c.execute(text("SELECT mode, status, task_id FROM migration_runs WHERE id = :r"),
                        {"r": body["run_id"]}).fetchone()
    engine.dispose()
    assert tuple(row) == ("s4_dry_run", "queued", "task-1")
