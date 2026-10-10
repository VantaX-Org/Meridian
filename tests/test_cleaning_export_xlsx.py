"""Cleaning queue export endpoint (/cleaning/export/{export_format}).

I9: build a real workbook through the route function and reload with
openpyxl, and confirm the batch_id filter added for N1 is wired into the
query's WHERE clause.
"""

import asyncio
import io
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock

import openpyxl

from api.routes.cleaning import export_cleaning_data


def _row(record_key: str, object_type: str, after: dict) -> SimpleNamespace:
    return SimpleNamespace(_mapping={
        "record_key": record_key,
        "object_type": object_type,
        "record_data_before": None,
        "record_data_after": after,
    })


def test_export_cleaning_xlsx_builds_real_workbook_and_filters_by_batch():
    tenant = SimpleNamespace(id="00000000-0000-0000-0000-000000000001", name="Acme Corp")
    rows = [_row("PARTNER=4711", "business_partner", {"NAME1": "Acme Trading"})]

    db = AsyncMock()
    captured_params = {}

    async def fake_execute(stmt, params=None):
        captured_params.update(params or {})
        return MagicMock(fetchall=lambda: rows)

    db.execute.side_effect = fake_execute

    response = asyncio.run(
        export_cleaning_data(
            export_format="xlsx",
            status="approved",
            object_type=None,
            batch_id="11111111-1111-1111-1111-111111111111",
            db=db,
            tenant=tenant,
        )
    )

    # N1: the batch_id filter must reach the query so a batch-scoped export
    # doesn't silently pull in every other batch's approved items.
    assert captured_params.get("bid") == "11111111-1111-1111-1111-111111111111"

    disposition = response.headers["content-disposition"]
    filename = disposition.split("filename=", 1)[1].strip('"')
    assert filename.startswith("meridian-")

    wb = openpyxl.load_workbook(io.BytesIO(response.body))
    sheet_name = next(s for s in wb.sheetnames if s != "Cover")
    assert "business_partner" in sheet_name
    ws = wb[sheet_name]
    headers = [c.value for c in ws[1]]
    assert "PARTNER" in headers
    assert "NAME_ORG1" in headers
