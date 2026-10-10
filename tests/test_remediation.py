"""Remediation batches: proposed values, current-value lookup, export layouts."""

import asyncio
import io
import uuid
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock

import openpyxl
import pandas as pd

from api.routes.remediation import export_batch
from api.services import remediation
from checks.frames import TableFrames
from sap.ddic import get_dictionary


def _frames():
    df = pd.DataFrame({"LFA1.LIFNR": ["0000100001", "0000100002"], "LFA1.LAND1": ["ZZ", None]})
    return TableFrames({"LFA1": df}, get_dictionary("s4hana"))


def test_build_items_reads_current_and_rule_fix(monkeypatch):
    monkeypatch.setattr(remediation, "module_rules",
                        lambda m: {"AP_T": {"id": "AP_T", "check_class": "null_check", "field": "LFA1.LAND1",
                                         "fix_value": {"__blank__": "DE"}}})
    issues = [{"issue_id": "i1", "scope": "upload", "module": "accounts_payable", "check_id": "AP_T",
               "record_key": "LIFNR=0000100002", "grain": "LFA1", "field": "LFA1.LAND1"},
              {"issue_id": "i2", "scope": "upload", "module": "accounts_payable", "check_id": "AP_T",
               "record_key": "LIFNR=0000100001", "grain": "LFA1", "field": "LFA1.LAND1"}]
    a, b = remediation.build_items(issues, _frames())
    assert (a["current_value"], a["proposed_value"], a["proposal_source"]) == (None, "DE", "rule")
    assert a["confidence"] == "medium" and b["confidence"] is None
    assert (b["current_value"], b["proposed_value"], b["proposal_source"]) == ("ZZ", None, "manual")


def test_exports_skip_blank_proposals():
    items = [
        {"module": "accounts_payable", "check_id": "AP005", "field": "LFB1.ZTERM", "proposed_value": "0001",
         "current_value": "", "record_key": "LIFNR=1|BUKRS=1000"},
        {"module": "accounts_payable", "check_id": "AP006", "field": "LFB1.AKONT", "proposed_value": "160000",
         "current_value": "1", "record_key": "LIFNR=1|BUKRS=1000"},
        {"module": "accounts_payable", "check_id": "AP007", "field": "LFA1.LAND1", "proposed_value": None,
         "current_value": "ZZ", "record_key": "LIFNR=2"},
    ]
    sheets = remediation.cockpit_sheets(items, get_dictionary("s4hana"))
    assert list(sheets) == ["LFB1"]
    assert sheets["LFB1"].to_dict("records") == [{"LIFNR": "1", "BUKRS": "1000", "ZTERM": "0001", "AKONT": "160000"}]
    assert list(remediation.cockpit_csv(items)["TABLE"]) == ["LFB1"]
    mc = remediation.mass_change(items)
    assert list(mc["NEW_VALUE"]) == ["0001", "160000"]
    assert mc.iloc[1][["TCODE", "TABLE", "FIELD", "OLD_VALUE"]].tolist() == ["FK02", "LFB1", "AKONT", "1"]


def test_export_batch_cockpit_xlsx_builds_real_workbook():
    """I9: build a real workbook through the remediation cockpit export route and
    reload it with openpyxl — sheet/column names, N3's narrowed SAP-reimport guard
    on a negative value, and the shared "meridian-" filename stamp."""
    batch_id = uuid.uuid4()
    tenant = SimpleNamespace(id=uuid.uuid4(), name="Acme Corp")

    batch_row = SimpleNamespace(_mapping={"id": str(batch_id), "status": "approved"})
    item_row = SimpleNamespace(_mapping={
        "module": "accounts_payable", "check_id": "AP005", "field": "LFB1.ZTERM",
        "proposed_value": "-5", "current_value": "0001", "record_key": "LIFNR=1|BUKRS=1000",
    })

    db = AsyncMock()
    db.execute.side_effect = [
        MagicMock(),  # RLS set_config
        MagicMock(fetchone=lambda: batch_row),  # _batch select
        MagicMock(fetchall=lambda: [item_row]),  # _items select
        MagicMock(),  # UPDATE remediation_batches ... status = 'exported'
        MagicMock(),  # _event insert
    ]
    db.commit = AsyncMock()

    response = asyncio.run(
        export_batch(batch_id=batch_id, request=None, format="cockpit_xlsx", db=db, tenant=tenant)
    )

    disposition = response.headers["content-disposition"]
    filename = disposition.split("filename=", 1)[1].strip('"')
    assert filename.startswith("meridian-")

    async def _collect() -> bytes:
        chunks = [c async for c in response.body_iterator]
        return b"".join(chunks)

    wb = openpyxl.load_workbook(io.BytesIO(asyncio.run(_collect())))
    assert "LFB1" in wb.sheetnames
    ws = wb["LFB1"]
    headers = [c.value for c in ws[1]]
    assert headers == ["LIFNR", "BUKRS", "ZTERM"]
    # N3: a leading "-" on a proposed value must survive unguarded for SAP reimport.
    assert ws.cell(row=2, column=3).value == "-5"
