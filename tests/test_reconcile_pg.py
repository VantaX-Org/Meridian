"""GET /api/v1/rules/{check_id}/reconcile and the reconciliation workbook against a real Postgres (seed: tests/m2_seed.py)."""
from __future__ import annotations

import io
import os

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine, text

from tests.m2_client import USER, headers, make_client
from tests.m2_seed import R1, R1B, R2, R3, seed_m2

pytestmark = pytest.mark.skipif(not os.environ.get("MERIDIAN_TEST_DB_URL"), reason="MERIDIAN_TEST_DB_URL not set")


@pytest.fixture(scope="module")
def owner():
    eng = create_engine(os.environ["MERIDIAN_TEST_DB_URL"])
    yield eng
    eng.dispose()


async def start(monkeypatch, tid, sub=USER):
    deps, client = make_client(monkeypatch, tid, sub)
    await deps.engine.dispose()
    return client


@pytest.mark.anyio
async def test_reconcile_a_rule_of_a_run(owner, monkeypatch):
    s = seed_m2(owner)
    client = await start(monkeypatch, s["tid"])
    async with client:
        r = await client.get("/api/v1/rules/MM014/reconcile", params={"run": s["run"]}, headers=headers("viewer"))
        assert r.status_code == 200, r.text
        b = r.json()
        assert (b["check_id"], b["module"], b["label"], b["run_id"]) == ("MM014", "material_master", "Material master", s["run"])
        assert (b["check_class"], b["field"], b["system"], b["transaction"], b["se16"]) == (
            "null_check", "MARA.NTGEW", "abap", "SE16N", True)
        # the counts are the run's stored finding, not recomputed
        assert (b["population"], b["failing"]) == (61816, 3640)
        assert b["grain"] == {"table": "MARA", "key_fields": ["MATNR", "WERKS"]}
        assert [(x["table"], x["field"], x["operator"]) for x in b["population_selection"]] == [("MARA", "LVORM", "initial")]
        assert [(x["table"], x["field"], x["operator"]) for x in b["failing_selection"]] == [("MARA", "NTGEW", "initial")]
        # five keys at most, in key order, ALPHA-formatted with the stored key alongside
        assert [k["raw"] for k in b["sample_keys"]] == sorted([R1, R1B, R2, R3])
        assert b["sample_keys"][0]["display"] == "MATNR 101, WERKS 1000"

        # a rule with a pattern says what SE16 cannot do
        with owner.begin() as c:
            c.execute(text("INSERT INTO findings (version_id, tenant_id, module, check_id, severity, dimension, affected_count, "
                           "total_count, details) VALUES (:v, :t, 'material_master', 'MM002', 'medium', 'validity', 4, 900, "
                           "CAST('{}' AS jsonb))"), {"v": s["run"], "t": s["tid"]})
        mm2 = (await client.get("/api/v1/rules/MM002/reconcile", params={"run": s["run"]}, headers=headers())).json()
        assert mm2["se16"] is False and mm2["conditions"][0]["kind"] == "pattern" and mm2["sample_keys"] == []

        # 404 sentences: an unknown run, a malformed run id, an unknown rule, a rule the run did not measure
        for params, rule, says in (({"run": s["run0"][:-1] + ("0" if s["run0"][-1] != "0" else "1")}, "MM014", "was not found"),
                                   ({"run": "not-a-run"}, "MM014", "was not found"),
                                   ({"run": s["run"]}, "NOPE999", "is not a rule of this tenant"),
                                   ({"run": s["run"]}, "MM200A", "is not a rule of this tenant")):
            bad = await client.get(f"/api/v1/rules/{rule}/reconcile", params=params, headers=headers())
            assert bad.status_code == 404 and says in bad.json()["detail"], bad.text
        miss = await client.get("/api/v1/rules/MM020/reconcile", params={"run": s["run"]}, headers=headers())
        assert miss.status_code == 404 and miss.json()["detail"] == f"Rule MM020 was not measured in run {s['run']}."
        assert (await client.get("/api/v1/rules/MM014/reconcile", headers=headers())).status_code == 422  # run is required


@pytest.mark.anyio
async def test_reconcile_is_tenant_scoped(owner, monkeypatch):
    s = seed_m2(owner)
    from tests.landscape_seed import seed
    other = seed(owner, with_run=False)
    client = await start(monkeypatch, other["tid"])
    async with client:
        r = await client.get("/api/v1/rules/MM014/reconcile", params={"run": s["run"]}, headers=headers())
        assert r.status_code == 404 and r.json()["detail"] == f"Run {s['run']} was not found."
        w = await client.get(f"/api/v1/runs/{s['run']}/reconciliation.xlsx", headers=headers())
        assert w.status_code == 404


@pytest.mark.anyio
async def test_reconciliation_workbook(owner, monkeypatch):
    s = seed_m2(owner)
    with owner.begin() as c:  # a key that looks like a formula must reach the sheet as text
        c.execute(text("INSERT INTO finding_records (tenant_id, version_id, check_id, module, grain, record_key) "
                       "VALUES (:t, :v, 'MM014', 'material_master', 'MARA', '@SUM(1+1)')"), {"t": s["tid"], "v": s["run"]})
    client = await start(monkeypatch, s["tid"])
    async with client:
        assert (await client.get(f"/api/v1/runs/{s['run']}/reconciliation.xlsx", headers=headers("viewer"))).status_code == 403
        r = await client.get(f"/api/v1/runs/{s['run']}/reconciliation.xlsx", params={"top": 2}, headers=headers("analyst"))
        assert r.status_code == 200, r.text
        assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
        wb = load_workbook(io.BytesIO(r.content))
        ranked = [x[0] for x in (owner.connect().execute(text(
            "SELECT check_id FROM findings WHERE tenant_id = :t AND version_id = :v AND affected_count > 0 "
            "ORDER BY affected_count DESC, check_id LIMIT 2"), {"t": s["tid"], "v": s["run"]})).fetchall()]
        assert wb.sheetnames == ["Cover", "Summary", *ranked]
        summary = list(wb["Summary"].iter_rows(values_only=True))
        assert summary[0] == ("Rule", "Object", "Table (grain)", "SE16 selection", "SE16 alone", "Population (Meridian)",
                              "Failing (Meridian)", "SE16 count", "Match", "Note")
        assert [row[0] for row in summary[1:]] == ranked
        mm14 = next(row for row in summary[1:] if row[0] == "MM014")
        assert mm14[5:9] == (61816, 3640, None, None)  # the team fills in the SE16 count and the match
        items = {row[0]: row[1] for row in wb["MM014"].iter_rows(values_only=True)}
        assert items["Failing (Meridian)"] == "3,640" and items["Transaction"] == "SE16N"
        keys = [row[1] for row in wb["MM014"].iter_rows(values_only=True) if row[0] == "Sample failing key"]
        assert keys[0] == "'@SUM(1+1)  (stored as @SUM(1+1))" and len(keys) == 5
        audit = owner.connect().execute(text(
            "SELECT action, entity_id FROM audit_log WHERE tenant_id = :t AND action = 'export.reconciliation'"),
            {"t": s["tid"]}).fetchall()
        assert [(a, e) for a, e in audit] == [("export.reconciliation", s["run"])]
