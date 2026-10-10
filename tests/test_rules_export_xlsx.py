"""Rules export endpoints (T10): /rules/export, /rules/{rule_id}/history/export.

No DB fixture — checks signatures and route ordering (export must not be
shadowed by the /{rule_id} catch-all route).
"""

import asyncio
import inspect
import io
import uuid
from types import SimpleNamespace
from unittest.mock import AsyncMock

import openpyxl

import api.routes.rules as rules_module
from api.routes.rules import export_rule_history, export_rules, router


def test_export_rules_signature() -> None:
    sig = inspect.signature(export_rules)
    assert "format" in sig.parameters
    assert sig.parameters["format"].default.default == "xlsx"
    for name in ("category", "module", "severity", "enabled", "search", "source"):
        assert name in sig.parameters


def test_export_rules_builds_real_workbook() -> None:
    """I9: build a real workbook through export_rules and reload with openpyxl —
    sheet/column names, the percent-formatted last_pass_rate, and the cover's
    formula-injection guard on tenant name."""
    rule = {
        "id": "RULE-1",
        "name": "Material description not blank",
        "module": "material_master",
        "category": "completeness",
        "severity": "high",
        "enabled": True,
        "source": "standard",
        "last_pass_rate": 88.25,
        "last_run_at": "2026-02-01T00:00:00",
        "description": "Checks MAKT-MAKTX is populated",
    }

    async def fake_list_rules(**kwargs):
        return {"rules": [rule]}

    db = AsyncMock()
    tenant = SimpleNamespace(id=uuid.uuid4(), name="=1+1")

    original = rules_module.list_rules
    rules_module.list_rules = fake_list_rules
    try:
        response = asyncio.run(export_rules(db=db, tenant=tenant))
    finally:
        rules_module.list_rules = original

    wb = openpyxl.load_workbook(io.BytesIO(response.body))
    assert wb.sheetnames == ["Cover", "Rules"]
    ws = wb["Rules"]
    headers = [c.value for c in ws[1]]
    assert headers == [
        "ID", "Name", "Module", "Category", "Severity", "Enabled", "Source",
        "Last pass rate", "Last run", "Description",
    ]
    pass_rate_cell = ws.cell(row=2, column=headers.index("Last pass rate") + 1).value
    assert abs(pass_rate_cell - 0.8825) < 1e-6

    cover = wb["Cover"]
    assert cover.cell(row=6, column=2).value.startswith("'=")


def test_export_rule_history_signature() -> None:
    sig = inspect.signature(export_rule_history)
    assert "format" in sig.parameters
    assert "rule_id" in sig.parameters


def test_export_route_declared_before_rule_id_catchall() -> None:
    paths = [r.path for r in router.routes]
    assert paths.index("/api/v1/rules/export") < paths.index("/api/v1/rules/{rule_id}")
