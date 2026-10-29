"""Tests for api.services.branded_xlsx — see ui-brand-exports-addendum.md T7."""

from datetime import datetime, timezone

import openpyxl
import pytest

from api.services.branded_xlsx import (
    ColumnSpec,
    SheetSpec,
    build_workbook,
    guard_formula_cell,
    guard_frame,
    guard_sap_reimport_cell,
    xlsx_filename,
)


def _load(data: bytes) -> openpyxl.Workbook:
    import io

    return openpyxl.load_workbook(io.BytesIO(data))


def _basic_sheet(rows=None) -> SheetSpec:
    return SheetSpec(
        title="Findings",
        columns=[
            ColumnSpec(key="check_id", header="Check ID", kind="mono"),
            ColumnSpec(key="affected", header="Affected", kind="int"),
            ColumnSpec(key="pass_rate", header="Pass rate", kind="pct"),
            ColumnSpec(key="note", header="Note", kind="text"),
        ],
        rows=rows if rows is not None else [
            {"check_id": "BP-001", "affected": 12, "pass_rate": 87.5, "note": "ok"},
        ],
    )


def test_cover_sheet_first_and_active_with_metadata():
    wb = _load(
        build_workbook(
            tenant_name="Acme Corp",
            run_label="Nightly run",
            run_id="abc-123",
            title="Findings export",
            sheets=[_basic_sheet()],
        )
    )
    assert wb.sheetnames[0] == "Cover"
    assert wb.active.title == "Cover"
    ws = wb["Cover"]
    text_cells = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert "Meridian" in text_cells
    assert "Findings export" in text_cells
    assert "Acme Corp" in text_cells
    assert "Nightly run" in text_cells
    assert "abc-123" in text_cells
    assert "SAST" in text_cells


def test_data_sheet_freeze_panes_and_autofilter():
    wb = _load(build_workbook(tenant_name="T", run_label="R", run_id="1", title="X", sheets=[_basic_sheet()]))
    ws = wb["Findings"]
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref is not None


def test_header_fill_is_brand_accent():
    wb = _load(build_workbook(tenant_name="T", run_label="R", run_id="1", title="X", sheets=[_basic_sheet()]))
    ws = wb["Findings"]
    header_cell = ws.cell(row=1, column=1)
    assert "2563EB" in str(header_cell.fill.fgColor.rgb)


def test_number_formats_per_kind():
    wb = _load(build_workbook(tenant_name="T", run_label="R", run_id="1", title="X", sheets=[_basic_sheet()]))
    ws = wb["Findings"]
    # row 2: Check ID (mono), Affected (int), Pass rate (pct), Note (text)
    assert ws.cell(row=2, column=2).number_format == "#,##0"
    assert ws.cell(row=2, column=3).number_format == "0.0%"


def test_formula_injection_guard_on_text_column():
    guarded = guard_formula_cell("=1+1")
    assert guarded == "'=1+1"
    wb = _load(
        build_workbook(
            tenant_name="T",
            run_label="R",
            run_id="1",
            title="X",
            sheets=[_basic_sheet(rows=[{"check_id": "BP-001", "affected": 1, "pass_rate": 1, "note": "=1+1"}])],
        )
    )
    ws = wb["Findings"]
    assert ws.cell(row=2, column=4).value == "'=1+1"


def test_cover_sheet_note_has_no_mid_string_apostrophe():
    """N5: a sheet note starting with '=' must not leave a literal apostrophe
    in the middle of the cover's "Sheets" sentence — only the cell's actual
    leading character matters for the formula guard."""
    wb = _load(
        build_workbook(
            tenant_name="T",
            run_label="R",
            run_id="1",
            title="X",
            sheets=[SheetSpec(title="Findings", columns=[], rows=[], note="=1+1 rows truncated")],
        )
    )
    ws = wb["Cover"]
    text_cells = [str(c.value) for row in ws.iter_rows() for c in row if c.value]
    sheet_row = next(t for t in text_cells if "Findings" in t and "row" in t)
    assert sheet_row == "Findings — 0 rows. =1+1 rows truncated"
    assert "'" not in sheet_row


def test_sap_reimport_guard_leaves_negative_values_untouched():
    """N3: a SAP-reimport file (kind='raw') must not prefix a leading
    apostrophe onto a negative quantity or a '-'-prefixed SAP code, or the
    reimport receives corrupted data. It must still guard '=' / '@' / tab /
    CR, which do evaluate as formulas."""
    assert guard_sap_reimport_cell("-5") == "-5"
    assert guard_sap_reimport_cell("-X100") == "-X100"
    assert guard_sap_reimport_cell("+5") == "+5"
    assert guard_sap_reimport_cell("=1+1") == "'=1+1"
    assert guard_sap_reimport_cell("@SUM(A1)") == "'@SUM(A1)"
    assert guard_sap_reimport_cell("\tx") == "'\tx"

    wb = _load(
        build_workbook(
            tenant_name="T",
            run_label="R",
            run_id="1",
            title="X",
            sheets=[
                SheetSpec(
                    title="Cockpit",
                    columns=[ColumnSpec(key="qty", header="Qty", kind="raw")],
                    rows=[{"qty": "-5"}],
                )
            ],
        )
    )
    ws = wb["Cockpit"]
    assert ws.cell(row=2, column=1).value == "-5"


def test_row_cap_truncates_and_notes_on_cover():
    rows = [{"check_id": f"BP-{i}", "affected": i, "pass_rate": 1, "note": "x"} for i in range(100_001)]
    wb = _load(
        build_workbook(
            tenant_name="T", run_label="R", run_id="1", title="X", sheets=[_basic_sheet(rows=rows)]
        )
    )
    ws = wb["Findings"]
    assert ws.max_row == 100_001  # header + 100,000 data rows
    cover_text = " ".join(str(c.value) for row in wb["Cover"].iter_rows() for c in row if c.value)
    assert "Truncated" in cover_text


def test_sheet_titles_truncated_to_31_chars_and_unique():
    long_title = "This Sheet Title Is Way Too Long For Excel"
    wb = _load(
        build_workbook(
            tenant_name="T",
            run_label="R",
            run_id="1",
            title="X",
            sheets=[
                SheetSpec(title=long_title, columns=_basic_sheet().columns, rows=[]),
                SheetSpec(title=long_title, columns=_basic_sheet().columns, rows=[]),
            ],
        )
    )
    data_titles = [t for t in wb.sheetnames if t != "Cover"]
    assert all(len(t) <= 31 for t in data_titles)
    assert len(set(data_titles)) == 2


def test_generated_at_utc_renders_two_hours_later_in_sast():
    generated = datetime(2025, 6, 1, 10, 0, tzinfo=timezone.utc)
    wb = _load(
        build_workbook(
            tenant_name="T",
            run_label="R",
            run_id="1",
            title="X",
            sheets=[_basic_sheet()],
            generated_at=generated,
        )
    )
    cover_text = " ".join(str(c.value) for row in wb["Cover"].iter_rows() for c in row if c.value)
    assert "12:00" in cover_text  # UTC+2 SAST


def test_xlsx_filename_format():
    name = xlsx_filename("findings", "Nightly Run", datetime(2025, 6, 1, 10, 0, tzinfo=timezone.utc))
    assert name.startswith("meridian-findings-nightly-run-")
    assert name.endswith("-SAST.xlsx")


def test_guard_frame_is_the_shared_export_guard():
    """One guard for every DataFrame export writer: full prefixes on report columns
    (leading spaces included), the narrowed SAP guard on SAP-mapped columns."""
    import pandas as pd

    df = pd.DataFrame({"note": ["=1+1", " =cmd", "-x", "\tx", "\rx", "ok"],
                       "MENGE": ["-5", "+5", "=1+1", "@SUM(A1)", " -1", "10"],
                       "qty": [1, -2, 3, 4, 5, 6]})
    out = guard_frame(df, sap_columns=["MENGE"])
    assert out["note"].tolist() == ["'=1+1", "' =cmd", "'-x", "'\tx", "'\rx", "ok"]
    assert out["MENGE"].tolist() == ["-5", "+5", "'=1+1", "'@SUM(A1)", " -1", "10"]
    assert out["qty"].tolist() == [1, -2, 3, 4, 5, 6]
    assert df["note"][0] == "=1+1"  # input untouched


def test_guard_frame_strips_control_characters_so_the_workbook_writes():
    import io

    import pandas as pd
    from openpyxl import load_workbook

    df = pd.DataFrame({"note": ["a\x0bb", "c\x00d", "ok"], "qty": [1, 2, 3]})
    out = guard_frame(df)
    assert out["note"].tolist() == ["ab", "cd", "ok"]
    assert df["note"][0] == "a\x0bb"  # input untouched
    buf = io.BytesIO()
    out.to_excel(buf, index=False)
    ws = load_workbook(io.BytesIO(buf.getvalue())).active
    assert [ws.cell(row=r, column=1).value for r in (2, 3, 4)] == ["ab", "cd", "ok"]


def test_cover_band_mark_meta_hyperlinks_and_tabs():
    wb = _load(build_workbook(tenant_name="T", run_label="R", run_id="rid-1", title="Findings export",
                              sheets=[_basic_sheet()]))
    ws = wb["Cover"]
    assert ws.column_dimensions["A"].width == 4
    assert all("F4F6F8" in str(ws.cell(row=r, column=1).fill.fgColor.rgb) for r in (1, 20, 40))
    assert len(ws._images) == 1 and ws._images[0].anchor._from.row == 1  # B2
    assert ws["B5"].value == "Findings export" and ws["B5"].font.size == 20 and ws["B5"].font.bold
    labels = {ws.cell(row=r, column=2).value: ws.cell(row=r, column=3) for r in range(8, 12)}
    assert labels["Run ID"].value == "rid-1" and labels["Run ID"].font.name == "Consolas"
    links = [c.hyperlink.location or c.hyperlink.target for row in ws.iter_rows() for c in row if c.hyperlink]
    assert any("Findings" in str(link) for link in links)
    assert "stays inside your network" in " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert "1E3A8A" in str(ws.sheet_properties.tabColor.rgb)
    assert "2563EB" in str(wb["Findings"].sheet_properties.tabColor.rgb)


def test_data_sheet_wash_alternation_and_status_tint():
    sheet = SheetSpec(
        title="S",
        columns=[ColumnSpec("v", "V", fill_by_value={"bad": "FDECEC"})],
        rows=[{"v": "a"}, {"v": "b"}, {"v": "c"}, {"v": "bad"}],
    )
    ws = _load(build_workbook(tenant_name="T", run_label="R", run_id="1", title="X", sheets=[sheet]))["S"]
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref
    assert "F4F6F8" not in str(ws["A2"].fill.fgColor.rgb) and "F4F6F8" in str(ws["A3"].fill.fgColor.rgb)
    assert "FDECEC" in str(ws["A5"].fill.fgColor.rgb) and "1A2233" in str(ws["A5"].font.color.rgb)


def test_cover_sheet_list_guards_formula_titles_and_escapes_apostrophes():
    from api.services import branded_xlsx as bx

    wb = _load(build_workbook(tenant_name="T", run_label="R", run_id="r", title="X",
                              sheets=[SheetSpec(title="=SUM(1)", columns=[ColumnSpec("v", "V")], rows=[]),
                                      SheetSpec(title="It's", columns=[ColumnSpec("v", "V")], rows=[])]))
    cells = [c for row in wb["Cover"].iter_rows() for c in row if c.hyperlink]
    assert cells and all(c.data_type != "f" for c in cells)
    assert any(c.value == "'=SUM(1)" for c in cells)
    assert any(c.hyperlink.target == "#'It''s'!A1" for c in cells)
    assert bx._MARK_PATH.name == "mark-light.png"  # anchored on the white B2 cell


@pytest.mark.parametrize("kind", ["text", "mono", "raw"])
def test_control_characters_in_sap_text_are_dropped_not_fatal(kind):
    """A SAP description with a vertical tab or NUL used to raise IllegalCharacterError and fail the whole export."""
    spec = SheetSpec(title="Records", columns=[ColumnSpec(key="name", header="Name", kind=kind)],
                     rows=[{"name": "Hex\x0bbolt\x00 M8"}])
    data = build_workbook(tenant_name="T", run_label=None, run_id="r", title="Fix sheet", sheets=[spec])
    assert _load(data)["Records"]["A2"].value == "Hexbolt M8"
