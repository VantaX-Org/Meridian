"""PowerPoint export of the L1-L5 process model (api/services/pptx_export.py)."""

from __future__ import annotations

from pptx import Presentation

from api.services.pptx_export import export_pptx
from sap.process_definitions import reference_document


def test_pptx_export_renders_a_slide_per_l1_and_l4() -> None:
    doc = reference_document()
    meta = {"name": "reference", "version_no": 0, "exported_at": "2026-01-01T00:00:00Z", "app_version": "test"}

    data = export_pptx(doc, meta)

    prs = Presentation(__import__("io").BytesIO(data))
    n_l1 = len(doc.l1)
    n_l4 = len(doc.all_l4())
    assert len(prs.slides) == 1 + n_l1 + n_l4  # title + per-L1 + per-L4
    # the first slide carries the model name; the L1 slides carry the top-level process names
    texts: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
    joined = "\n".join(texts)
    assert "reference" in joined
    for l1 in doc.l1:
        assert l1.name in joined


def test_pptx_export_is_nonempty_for_a_model_with_no_l1() -> None:
    from api.models.process_model import ProcessModelDocument

    data = export_pptx(ProcessModelDocument(l1=[]), {"name": "empty", "version_no": 1, "exported_at": "", "app_version": ""})
    assert len(data) > 0
    prs = Presentation(__import__("io").BytesIO(data))
    assert len(prs.slides) == 1  # title only
