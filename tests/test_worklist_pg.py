"""The steward work list (api/services/worklist.py, api/routes/work.py) against a real Postgres.

One run of a system, two checks on material master:
  MM001 critical  MARA.MEINS  MATNR=1 (auto_fix proposes EA), MATNR=2 (no proposal)
  MM002 medium    MARA.MTART  MATNR=2 (current XX, the rule's value fix says FERT)
The build runs as a non-superuser role, so RLS is enforced on the new tables.
"""
from __future__ import annotations

import json
import os
import uuid

import pandas as pd
import pytest
from httpx import ASGITransport, AsyncClient
from sqlalchemy import create_engine, text
from sqlalchemy.orm import Session

pytestmark = pytest.mark.skipif(not os.environ.get("MERIDIAN_TEST_DB_URL"), reason="MERIDIAN_TEST_DB_URL not set")

HEADERS = {"X-User-Role": "admin", "Authorization": "Bearer test-token"}
USER = "00000000-0000-0000-0000-0000000000a1"


@pytest.fixture(scope="module")
def engines():
    from urllib.parse import urlparse, urlunparse

    url = os.environ["MERIDIAN_TEST_DB_URL"]
    owner = create_engine(url)
    with owner.begin() as c:
        if c.execute(text("SELECT 1 FROM pg_roles WHERE rolname = 'meridian_work_app'")).scalar():
            c.execute(text("DROP OWNED BY meridian_work_app CASCADE"))
        c.execute(text("DROP ROLE IF EXISTS meridian_work_app"))
        c.execute(text("CREATE ROLE meridian_work_app LOGIN PASSWORD 'pw' NOSUPERUSER NOBYPASSRLS"))
        c.execute(text("GRANT USAGE, CREATE ON SCHEMA public TO meridian_work_app"))
        c.execute(text("GRANT ALL ON ALL TABLES IN SCHEMA public TO meridian_work_app"))
    u = urlparse(url)
    app = create_engine(urlunparse((u.scheme, f"meridian_work_app:pw@{u.hostname}:{u.port or 5432}",
                                    u.path, "", u.query, "")))
    yield owner, app
    app.dispose()
    with owner.begin() as c:
        c.execute(text("DROP OWNED BY meridian_work_app CASCADE"))
        c.execute(text("DROP ROLE meridian_work_app"))
    owner.dispose()


class _Frames:
    def __init__(self, frames):
        self.frames = frames


@pytest.fixture
def run(engines):
    owner, app = engines
    tid, vid, sid = str(uuid.uuid4()), str(uuid.uuid4()), str(uuid.uuid4())
    with owner.begin() as c:
        c.execute(text("INSERT INTO tenants (id, name) VALUES (:id, :id)"), {"id": tid})
        c.execute(text("INSERT INTO users (id, tenant_id, email, name, role) VALUES (:u, :t, :e, 'Steward', 'admin') "
                       "ON CONFLICT (id) DO NOTHING"), {"u": USER, "t": tid, "e": f"{tid[:8]}@example.com"})
        c.execute(text("INSERT INTO analysis_versions (id, tenant_id, status, run_at, metadata, dqs_summary) "
                       "VALUES (:v, :t, 'complete', now(), CAST(:m AS jsonb), CAST('{\"material_master\": {}}' AS jsonb))"),
                  {"v": vid, "t": tid, "m": json.dumps({"system_id": sid, "modules": ["material_master"]})})
        for check, sev, field, n, extra in (
            ("MM001", "critical", "MARA.MEINS", 2,
             {"record_fixes": [{"record_key": "MATNR=1", "proposed_value": "EA", "confidence": "high",
                                "fix_instruction": "Set the base unit"}]}),
            ("MM002", "medium", "MARA.MTART", 1,
             {"value_fix_map": {"XX": {"suggested_value": "FERT", "fix_instruction": "Use FERT"}}}),
        ):
            c.execute(text("""
                INSERT INTO findings (version_id, tenant_id, module, check_id, severity, dimension, affected_count,
                                      total_count, details, record_fixes, value_fix_map)
                VALUES (:v, :t, 'material_master', :c, :s, 'completeness', :n, 10, CAST(:d AS jsonb),
                        CAST(:rf AS jsonb), CAST(:vf AS jsonb))"""),
                      {"v": vid, "t": tid, "c": check, "s": sev, "n": n,
                       "d": json.dumps({"message": f"{field} is wrong", "field_checked": field}),
                       "rf": json.dumps(extra.get("record_fixes")), "vf": json.dumps(extra.get("value_fix_map"))})
        for check, key, values in (("MM001", "MATNR=1", {"MARA.MEINS": None}), ("MM001", "MATNR=2", {"MARA.MEINS": ""}),
                                   ("MM002", "MATNR=2", {"MARA.MTART": "XX"})):
            c.execute(text("INSERT INTO finding_records (tenant_id, version_id, check_id, module, grain, record_key, "
                           "field_values) VALUES (:t, :v, :c, 'material_master', 'MARA', :k, CAST(:fv AS jsonb))"),
                      {"t": tid, "v": vid, "c": check, "k": key, "fv": json.dumps(values)})
            c.execute(text("""
                INSERT INTO record_issues (tenant_id, scope, module, check_id, record_key, grain, severity, status,
                                           first_seen_version, last_seen_version)
                VALUES (:t, :s, 'material_master', :c, :k, 'MARA',
                        (SELECT severity FROM findings WHERE version_id = :v AND check_id = :c), 'open', :v, :v)"""),
                      {"t": tid, "s": sid, "c": check, "k": key, "v": vid})
    frames = _Frames({"MAKT": pd.DataFrame({"MAKT.MATNR": ["1", "1", "2"], "MAKT.MAKTX": ["Schraube", "Bolt", "Nut"],
                                            "MAKT.SPRAS": ["D", "E", "E"]})})
    from api.services import worklist
    with Session(app) as s:
        stats = worklist.build(s, tid, vid, frames)
    yield owner, tid, vid, stats


def test_build_ranks_names_and_counts(run):
    owner, tid, vid, stats = run
    assert stats["groups"] == 2 and stats["records"] == 2 and stats["named"] == 2
    with owner.begin() as c:
        groups = {r.check_id: r for r in c.execute(text("SELECT * FROM work_groups WHERE version_id = :v"), {"v": vid})}
        recs = {r.record_key: r for r in c.execute(text("SELECT * FROM work_records WHERE version_id = :v"), {"v": vid})}
    assert groups["MM001"].records == 2 and groups["MM001"].proposed == 1 and groups["MM001"].statuses == {"to_fix": 2}
    assert groups["MM002"].proposed == 1  # the value fix for XX
    assert recs["MATNR=1"].name == "Bolt"  # English description preferred
    assert recs["MATNR=2"].checks == ["MM001", "MM002"] and recs["MATNR=2"].issues == 2 and recs["MATNR=2"].sev_rank == 0
    assert recs["MATNR=2"].proposed == 1 and recs["MATNR=2"].status == "to_fix"


@pytest.fixture
def client_for(monkeypatch):
    from api import deps as api_deps
    from api.deps import Tenant, get_tenant
    from api.main import app

    def make(tid):
        monkeypatch.setattr("api.middleware.local_auth._load_jwt_secret", lambda: "test-secret")
        monkeypatch.setattr("api.middleware.local_auth.decode_access_token",
                            lambda token, secret: {"sub": USER, "email": "steward@example.com", "role": "admin"})
        monkeypatch.setenv("MERIDIAN_DEV_ROLE_HEADER", "1")
        monkeypatch.setitem(app.dependency_overrides, get_tenant, lambda: Tenant(uuid.UUID(tid), "T", []))
        return api_deps, AsyncClient(transport=ASGITransport(app=app), base_url="http://test")
    return make


@pytest.mark.anyio
async def test_worklist_api_lists_explains_and_decides(run, client_for):
    owner, tid, vid, _ = run
    deps, client = client_for(tid)
    await deps.engine.dispose()
    async with client:
        get = lambda path, **params: client.get(f"/api/v1/work{path}", params={"version_id": vid, **params}, headers=HEADERS)
        decide = lambda *items: client.post("/api/v1/work/decide", json={"version_id": vid, "items": list(items)},
                                            headers=HEADERS)

        s = (await get("/summary")).json()
        assert s["objects"][0]["module"] == "material_master" and s["totals"]["open"] == 2 and s["built"]

        g = (await get("/groups")).json()
        assert [x["check_id"] for x in g["groups"]] == ["MM001", "MM002"] and g["groups"][0]["open"] == 2

        first = (await get("/records", limit=1)).json()
        assert first["total"] == 2 and [i["record_key"] for i in first["items"]] == ["MATNR=1"] and first["next"]
        second = (await get("/records", limit=1, after=first["next"])).json()
        assert [i["record_key"] for i in second["items"]] == ["MATNR=2"] and second["next"] is None
        assert (await get("/records", check_id="MM002")).json()["total"] == 1

        panel = (await get("/records/material_master/MATNR=2")).json()
        by = {i["check_id"]: i for i in panel["issues"]}
        assert panel["name"] == "Nut" and by["MM001"]["suggestion"] is None and by["MM001"]["current_value"] == ""
        assert by["MM002"]["suggestion"] == {"value": "FERT", "source": "value_map", "confidence": None, "why": "Use FERT"}
        assert panel["can"]["fix"] and panel["can"]["sign_off"]

        # accept the proposal for one check: the record waits on its other check, so it is still to fix
        r = await decide({"module": "material_master", "record_key": "MATNR=2", "check_id": "MM002", "action": "accept"})
        assert r.status_code == 200, r.text
        batch = r.json()["draft_batch_id"]
        panel = (await get("/records/material_master/MATNR=2")).json()
        by = {i["check_id"]: i for i in panel["issues"]}
        assert by["MM002"]["status"] == "proposed" and by["MM002"]["decision"]["value"] == "FERT"
        assert panel["status"] == "to_fix"

        # edit a value for the other: the whole record now has a proposal
        r = await decide({"module": "material_master", "record_key": "MATNR=2", "check_id": "MM001", "action": "edit",
                          "value": "PC"})
        assert r.json()["draft_batch_id"] == batch  # one draft per steward
        assert (await get("/records", status="proposed")).json()["items"][0]["record_key"] == "MATNR=2"

        # undo takes both values back out
        await decide({"module": "material_master", "record_key": "MATNR=2", "action": "undo"})
        assert (await get("/records/material_master/MATNR=2")).json()["status"] == "to_fix"

        # a sign-off needs a reason
        bad = await decide({"module": "material_master", "record_key": "MATNR=1", "action": "not_an_issue"})
        assert bad.status_code == 400
        ok = await decide({"module": "material_master", "record_key": "MATNR=1", "action": "not_an_issue",
                           "note": "Service material, no unit"})
        assert ok.status_code == 200
        assert (await get("/records", status="not_an_issue")).json()["items"][0]["record_key"] == "MATNR=1"
        assert (await get("/records")).json()["total"] == 1  # open only

        # search by name and by key value
        assert [i["record_key"] for i in (await client.get("/api/v1/work/search", params={"q": "nu", "version_id": vid},
                                                           headers=HEADERS)).json()["items"]] == ["MATNR=2"]
        assert [i["record_key"] for i in (await client.get("/api/v1/work/search", params={"q": "002", "version_id": vid},
                                                           headers=HEADERS)).json()["items"]] == ["MATNR=2"]

        # accept-all: only the open records with a proposal (MATNR=1 is signed off, MATNR=2 has none for MM001)
        r = await client.post("/api/v1/work/groups/MM001/accept-all", params={"version_id": vid}, headers=HEADERS)
        assert r.status_code == 200 and r.json()["accepted"] == 0
        r = await client.post("/api/v1/work/groups/MM002/accept-all", params={"version_id": vid}, headers=HEADERS)
        assert r.json()["accepted"] == 1
        g = {x["check_id"]: x for x in (await get("/groups")).json()["groups"]}
        assert g["MM002"]["statuses"] == {"proposed": 1}

        # assign to me
        await decide({"module": "material_master", "record_key": "MATNR=2", "action": "assign", "assignee": "me"})
        assert (await get("/records", assigned="me")).json()["total"] == 1

    with owner.begin() as c:
        items = c.execute(text("SELECT check_id, record_key, proposed_value, proposal_source FROM remediation_items "
                               "WHERE batch_id = :b"), {"b": batch}).fetchall()
    assert [tuple(i) for i in items] == [("MM002", "MATNR=2", "FERT", "rule")]


@pytest.mark.anyio
async def test_fix_sheet_round_trip_mass_edit_and_approval(run, client_for, monkeypatch):
    import io

    from openpyxl import load_workbook

    owner, tid, vid, _ = run
    deps, client = client_for(tid)
    await deps.engine.dispose()
    async with client:
        sheet = await client.get("/api/v1/work/export", params={"version_id": vid}, headers=HEADERS)
        assert sheet.status_code == 200, sheet.text
        wb = load_workbook(io.BytesIO(sheet.content))
        ws = wb["Fix sheet"]
        header = [c.value for c in ws[1]]
        rows = {(r[header.index("Record")].value, r[header.index("Check")].value): r for r in ws.iter_rows(min_row=2)}
        assert set(rows) == {("MATNR=1", "MM001"), ("MATNR=2", "MM001"), ("MATNR=2", "MM002")}
        assert rows[("MATNR=1", "MM001")][header.index("Suggested value")].value == "EA"
        assert rows[("MATNR=2", "MM002")][header.index("Suggested value")].value == "FERT"
        rows[("MATNR=2", "MM001")][header.index("New value")].value = "KG"
        rows[("MATNR=1", "MM001")][header.index("Decision")].value = "Not an issue"  # no reason: skipped
        rows[("MATNR=2", "MM002")][header.index("Decision")].value = "fix"
        buf = io.BytesIO()
        wb.save(buf)

        bad = await client.post("/api/v1/work/import", params={"version_id": vid}, headers=HEADERS,
                                files={"file": ("fix.xlsx", b"not a zip", "application/octet-stream")})
        assert bad.status_code == 422
        r = await client.post("/api/v1/work/import", params={"version_id": vid}, headers=HEADERS,
                              files={"file": ("fix.xlsx", buf.getvalue(), "application/octet-stream")})
        assert r.status_code == 200, r.text
        assert r.json()["rows"] == 2 and len(r.json()["skipped"]) == 1
        batch = r.json()["draft_batch_id"]
        rec = (await client.get("/api/v1/work/records/material_master/MATNR=2", params={"version_id": vid},
                                headers=HEADERS)).json()
        assert rec["status"] == "proposed"
        assert {i["check_id"]: i["decision"]["value"] for i in rec["issues"]} == {"MM001": "KG", "MM002": "FERT"}

        # mass edit: one value for the whole group
        r = await client.post("/api/v1/work/groups/MM001/set-value", params={"version_id": vid}, headers=HEADERS,
                              json={"value": "ST"})
        assert r.json()["updated"] == 2

        # a second person approves the draft: the records follow
        monkeypatch.setattr("api.middleware.local_auth.decode_access_token",
                            lambda token, secret: {"sub": str(uuid.uuid4()), "email": "lead@example.com", "role": "admin"})
        r = await client.post(f"/api/v1/remediation/batches/{batch}/approve", headers=HEADERS)
        assert r.status_code == 200, r.text
        listed = (await client.get("/api/v1/work/records", params={"version_id": vid, "status": "approved"},
                                   headers=HEADERS)).json()
        assert {i["record_key"] for i in listed["items"]} == {"MATNR=1", "MATNR=2"}


@pytest.mark.anyio
async def test_records_filter_by_a_key_part(run, client_for):
    owner, tid, vid, _ = run
    deps, client = client_for(tid)
    await deps.engine.dispose()
    async with client:
        get = lambda **params: client.get("/api/v1/work/records", params={"version_id": vid, **params}, headers=HEADERS)
        assert [i["record_key"] for i in (await get(part="MATNR=0002")).json()["items"]] == ["MATNR=2"]
        assert (await get(part="WERKS=1000")).json()["total"] == 0
        assert (await get(part="nonsense")).status_code == 422


@pytest.mark.anyio
async def test_records_filter_blocking_every_status_and_source_system(run, client_for):
    owner, tid, vid, _ = run
    with owner.begin() as c:
        sid = c.execute(text("SELECT scope FROM work_records WHERE version_id = :v LIMIT 1"), {"v": vid}).scalar()
    deps, client = client_for(tid)
    await deps.engine.dispose()
    async with client:
        get = lambda **params: client.get("/api/v1/work/records", params={"version_id": vid, **params}, headers=HEADERS)
        assert (await get(severity="blocking")).json()["total"] == 2
        assert (await get(severity="low")).json()["total"] == 0
        assert (await get(status="all")).json()["total"] == 2
        assert (await get(system=sid)).json()["total"] == 2
        assert (await get(system="elsewhere")).json()["total"] == 0
        groups = (await client.get("/api/v1/work/groups", params={"version_id": vid, "severity": "blocking"},
                                   headers=HEADERS)).json()["groups"]
        assert [g["check_id"] for g in groups] == ["MM001"]


@pytest.mark.anyio
async def test_accept_all_keeps_entered_values_and_cleaning_suggestions_can_be_accepted(run, client_for):
    owner, tid, vid, _ = run
    with owner.begin() as c:
        c.execute(text("""INSERT INTO cleaning_queue (tenant_id, object_type, status, confidence, record_key,
                                                      record_data_after, version_id, check_id, field)
                          VALUES (:t, 'material_master', 'detected', 90, 'MATNR=2',
                                  CAST('{"MARA.MEINS": "PC"}' AS jsonb), :v, 'MM001', 'MARA.MEINS')"""),
                  {"t": tid, "v": vid})
    deps, client = client_for(tid)
    await deps.engine.dispose()
    async with client:
        decide = lambda *items: client.post("/api/v1/work/decide", json={"version_id": vid, "items": list(items)},
                                            headers=HEADERS)
        # the panel offers the cleaning engine's value, and accepting it works
        panel = (await client.get("/api/v1/work/records/material_master/MATNR=2", params={"version_id": vid},
                                  headers=HEADERS)).json()
        mm001 = next(i for i in panel["issues"] if i["check_id"] == "MM001")
        assert mm001["suggestion"]["source"] == "cleaning" and mm001["suggestion"]["value"] == "PC"
        r = await decide({"module": "material_master", "record_key": "MATNR=2", "check_id": "MM001", "action": "accept"})
        assert r.status_code == 200, r.text
        batch = r.json()["draft_batch_id"]
        # a value the steward entered survives accept-all for the issue
        await decide({"module": "material_master", "record_key": "MATNR=1", "check_id": "MM001", "action": "edit", "value": "KG"})
        await client.post("/api/v1/work/groups/MM001/accept-all", params={"version_id": vid}, headers=HEADERS)
    with owner.begin() as c:
        items = dict(c.execute(text("SELECT record_key, proposed_value || ':' || proposal_source FROM remediation_items "
                                    "WHERE batch_id = :b AND check_id = 'MM001'"), {"b": batch}).fetchall())
    assert items == {"MATNR=1": "KG:steward", "MATNR=2": "PC:cleaning"}


@pytest.mark.anyio
async def test_a_run_counts_as_built_only_once_its_build_finished(run, client_for):
    owner, tid, vid, _ = run
    deps, client = client_for(tid)
    await deps.engine.dispose()
    async with client:
        summary = lambda: client.get("/api/v1/work/summary", params={"version_id": vid}, headers=HEADERS)
        assert (await summary()).json()["built"] is True
        with owner.begin() as c:  # a build that stopped half way leaves rows but no completion marker
            c.execute(text("UPDATE analysis_versions SET metadata = metadata - 'worklist_built_at' WHERE id = :v"),
                      {"v": vid})
        assert (await summary()).json()["built"] is False


@pytest.mark.anyio
async def test_fix_sheet_sign_offs_need_the_approve_permission(run, client_for, monkeypatch):
    """Import is gated on apply; each sign-off row still needs approve (decide checks every action)."""
    import io

    from openpyxl import load_workbook

    owner, tid, vid, _ = run
    deps, client = client_for(tid)
    await deps.engine.dispose()
    real = __import__("api.routes.work", fromlist=["_allowed"])._allowed

    async def no_approve(db, tenant, request, role, action):
        return False if action == "approve" else await real(db, tenant, request, role, action)

    async with client:
        sheet = await client.get("/api/v1/work/export", params={"version_id": vid}, headers=HEADERS)
        wb = load_workbook(io.BytesIO(sheet.content))
        ws = wb["Fix sheet"]
        header = [c.value for c in ws[1]]
        row = next(r for r in ws.iter_rows(min_row=2) if r[header.index("Record")].value == "MATNR=1")
        row[header.index("Decision")].value = "Not an issue"
        row[header.index("Reason")].value = "Service material"
        buf = io.BytesIO()
        wb.save(buf)
        monkeypatch.setattr("api.routes.work._allowed", no_approve)
        r = await client.post("/api/v1/work/import", params={"version_id": vid}, headers=HEADERS,
                              files={"file": ("fix.xlsx", buf.getvalue(), "application/octet-stream")})
    assert r.status_code == 403 and "approve" in r.json()["detail"]
    with owner.begin() as c:
        assert c.execute(text("SELECT count(*) FROM record_issues WHERE tenant_id = :t AND status = 'accepted'"),
                         {"t": tid}).scalar() == 0


@pytest.mark.anyio
async def test_records_carry_age_and_suggested_value(run, client_for):
    """Each listed record says when the oldest of its failing issues was first seen (the queue's Age) and the first
    suggested value of its checks (the rule's fix for the record, else for its current value), read per page."""
    owner, tid, vid, _ = run
    with owner.begin() as c:
        c.execute(text("UPDATE record_issues SET first_seen_at = now() - CASE check_id WHEN 'MM001' "
                       "THEN interval '9 days' ELSE interval '30 days' END WHERE tenant_id = :t"), {"t": tid})
    deps, client = client_for(tid)
    await deps.engine.dispose()
    async with client:
        items = (await client.get("/api/v1/work/records", params={"version_id": vid}, headers=HEADERS)).json()["items"]
    by = {i["record_key"]: i for i in items}
    assert by["MATNR=1"]["suggested_value"] == "EA"  # MM001's fix for the record
    assert by["MATNR=2"]["suggested_value"] == "FERT"  # MM001 has none for it: MM002's fix for the value XX
    now = pd.Timestamp.now(tz="UTC")
    assert (now - pd.Timestamp(by["MATNR=1"]["first_seen_at"])).days == 9
    assert (now - pd.Timestamp(by["MATNR=2"]["first_seen_at"])).days == 30  # the older of its two issues
