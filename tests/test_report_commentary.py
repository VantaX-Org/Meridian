"""The plain-speech commentary under every report section: the sentence builders, and the reports and
workbook that carry them. Everything here runs from fixtures, no database."""

from __future__ import annotations

import io
import re
import html

import pytest

from api.services import insights_aggregate as agg
from api.services import pdf_reports as pr
from api.services import report_commentary as rc
from tests import pdf_fixtures as fx

KW = {"tenant_name": fx.TENANT, "generated_at": fx.GENERATED, "scoring": {}}
T = {"target": 90.0, "warning": 75.0, "critical": 60.0}


def _text(template: str, ctx: dict) -> str:
    out = pr._env().get_template(template).render(**ctx)
    return html.unescape(re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", out)))


def _paras(template: str, ctx: dict) -> list[str]:
    out = pr._env().get_template(template).render(**ctx)
    return [html.unescape(re.sub(r"<[^>]+>", "", p)).strip()
            for p in re.findall(r'<p class="commentary">(.*?)</p>', out, flags=re.S)]


def board_data(with_previous: bool = True) -> dict:
    """What ``insights_aggregate.gather`` returns, built from the PDF fixtures."""
    findings = fx.FINDINGS2
    data = {
        "tenant_name": fx.TENANT, "version": fx.V2, "previous": fx.V1 if with_previous else None, "module": None,
        "findings": findings, "fixes_by_check": {"MM003": 4_000}, "previous_findings": fx.FINDINGS1 if with_previous else [],
        "scoring": {}, "waves": [{"name": "Wave 1", "modules": ["material_master", "fi_gl"]},
                                 {"name": "Wave 2", "modules": ["business_partner", "accounts_payable"]}],
        "fix_records": 4_000, "high_confidence": 2_500, "impact_rows": [], "currency": "ZAR",
        "clusters": agg.cluster_findings(findings, {"MM003": 4_000}), "system": fx.SYSTEM,
        "trend": {"runs": agg.trend_runs([fx.V1, fx.V2], {fx.V1["id"]: fx.FINDINGS1, fx.V2["id"]: fx.FINDINGS2}, {}, {}),
                  "annotations": []},
        "cleansing": {"steps": agg.cleansing_steps(agg.failing_records(findings), 4_000, 2_500, 900, 300)},
        # the workbook loader adds these on top of gather()
        "samples": [], "readiness": [], "tenant_id": "t1",
    }
    data["overview"] = agg.overview_from(data)
    return data


# ── sentence builders ────────────────────────────────────────────────────────


def test_overall_reads_the_score_against_the_bands():
    assert rc.overall(None, None) == "No module produced a score, so there is nothing to read against the target yet."
    s = rc.overall(72.3, -1.2, T, capped=True, critical=2)
    assert "72.3 is in the warning band of 90" in s and "17.7 more points" in s
    assert "1.2 points worse" in s and "2 critical checks fail and they cap the score" in s
    # a critical failure within tolerance does not cap: the scorer's verdict decides, not the count
    s = rc.overall(92.0, None, T, capped=False, critical=1)
    assert "1 critical check fails, within the tolerance that would cap the score" in s and "fastest" not in s
    s = rc.overall(93.0, 0.0, T)
    assert "meets the target" in s and "unchanged" in s and "critical" not in s
    assert "paying off" in rc.overall(80.0, 3.0, T)


def test_dimensions_name_the_weakest_and_what_it_means():
    dims = {"completeness": 61.0, "accuracy": 88.0, "consistency": None, "timeliness": None, "uniqueness": 99.0, "validity": 70.0}
    s = rc.dimensions(dims, {"completeness": 58.0, "accuracy": 88.0, "uniqueness": 99.0, "validity": 74.0}, T)
    assert s.startswith("Completeness is the weakest dimension at 61.0 and uniqueness the strongest at 99.0.")
    assert "required fields have been left empty" in s and "3 of 4 measured dimensions" in s
    assert "Consistency and timeliness were not measured, so no claim is made about them." in s
    assert "validity moved most, down 4.0 points" in s
    assert rc.dimensions({d: None for d in rc.DIMENSIONS}).startswith("No dimension was measured")
    one = rc.dimensions({"completeness": 95.0, "accuracy": None, "consistency": None, "timeliness": None,
                         "uniqueness": None, "validity": None}, thresholds=T)
    assert "Every measured dimension meets the target" in one and "about them" in one


def test_rule_names_are_shortened_inside_sentences():
    assert rc._short("Material type is mandatory - determines procurement and storage behaviour") == "Material type is mandatory"
    assert rc._short(None, "MM003") == "MM003"
    assert rc._short("x" * 90).endswith("…")


def test_fix_first_and_next_steps():
    rows = [{"check_id": "BP001", "message": "Partner has no name", "module": "business_partner", "affected": 3_120,
             "severity": "critical", "field": "BUT000.NAME_ORG1", "why": "A partner without a name cannot be invoiced."},
            {"check_id": "MM003", "message": "Material type is mandatory - determines storage", "module": "material_master",
             "affected": 500, "severity": "high"}]
    s = rc.fix_first(rows, capped=True)
    assert s.startswith("The first item alone, Partner has no name in Business partner, carries")
    assert "on BUT000.NAME_ORG1" in s and "A partner without a name cannot be invoiced." in s
    assert "1 of the 2 is critical" in s and s.endswith("Work the list top down: each item is worth more than everything below it.")
    assert "1 of the 2 is critical and comes first whatever the count." in rc.fix_first(rows)
    steps = rc.next_steps(fix_first_rows=rows, critical=1, capped=True, errored=2, score=72.3, thresholds=T)
    assert steps[0] == "Fix Partner has no name in Business partner first: 3,120 records, the largest impact in the run."
    assert steps[1] == "Clear the 1 critical failure to lift the cap on the score."
    assert "2 checks that could not run" in steps[2] and steps[3].endswith("17.7 points away.")
    assert rc.next_steps(critical=2) == ["Clear the 2 critical failures: each one marks records that cannot be used."]
    assert rc.next_steps() == []
    assert rc.fix_first([]) == ""


def test_severity_mix_and_critical_findings():
    assert "rare" in rc.severity_mix({"critical": 0, "high": 0, "medium": 0, "low": 0}, 10)
    s = rc.severity_mix({"critical": 0, "high": 2, "medium": 5, "low": 4}, 20)
    assert s.startswith("11 of 20 checks found failing records.") and "not capped" in s and "Most of the failures are medium or low" in s
    assert rc.critical_findings([]).startswith("No critical check fails")
    s = rc.critical_findings([{"module": "a", "affected": 10}, {"module": "a", "affected": 5}, {"module": "b", "affected": 1}])
    assert "3 critical checks fail across 2 modules, on 16 records in total. A carries the most, 2." in s
    tie = rc.critical_findings([{"module": "a", "affected": 1}, {"module": "b", "affected": 1}])
    assert "carries the most" not in tie


def test_clusters_blockers_and_waves():
    clusters = [{"module": "material_master", "field": "MTART", "impact": 70, "rules": ["MM003"], "auto_fix_pct": 35.0},
                {"module": "fi_gl", "field": None, "impact": 20, "rules": ["GL1", "GL2"]},
                {"module": "fi_gl", "field": "X", "impact": 5, "rules": ["GL3"]},
                {"module": "fi_gl", "field": "Y", "impact": 5, "rules": ["GL4"]}]
    s = rc.clusters(clusters)
    assert s.startswith("Material master MTART is the largest root cause, 70% of the impact across the clusters shown, through 1 rule.")
    assert "top three clusters together hold 95%" in s and "35% of the top cluster's records already have a proposed value" in s
    assert rc.clusters([]) == ""
    assert rc.blockers([], 0).startswith("No blocking rule fails")
    waves = [{"name": "Wave 1", "modules": [{"name": "material_master", "score": 84.1, "blockers": 1}]},
             {"name": "Wave 2", "modules": [{"name": "fi_gl", "score": 93.4, "blockers": 0}]}, {"name": "Wave 3", "modules": []}]
    s = rc.blockers([{"check_id": "MM003", "name": "Material type is mandatory - x", "module": "material_master", "affected": 11_420}], 1, waves)
    assert s == ("1 blocking rule fails. The largest is Material type is mandatory in Material master, failing on 11,420 records. "
                 "Every blocker shown sits in Material master, so one team owns the whole list. Wave 1 cannot go until these are cleared.")
    w = rc.waves(waves, T)
    assert "Wave 1 is held back: 1 of 1 module has blocking rules (Material master); Material master is below the 90 target." in w
    assert "Wave 2 is clear: no module in it has a blocking rule and every module meets the 90 target." in w
    assert "Wave 3 has no scored module." in w
    # between warning and target with no blocker is conditional, never clear
    cond = rc.waves([{"name": "Wave 1", "modules": [{"name": "fi_gl", "score": 80.0, "blockers": 0}]}], T)
    assert cond == "Wave 1 is not blocked, but FI GL scores below the 90 target, so it is conditional rather than ready."
    mixed = rc.waves([{"name": "W", "modules": [{"name": "fi_gl", "score": 80.0, "blockers": 0},
                                                 {"name": "hcm", "score": 50.0, "blockers": 0}]}], T)
    assert mixed == "W is held back: Hcm scores below 75; FI GL is below the 90 target."
    assert rc.waves([]).startswith("No migration wave is planned")


def test_module_paragraph():
    m = {"name": "material_master", "score": 84.1, "delta": -1.9, "dims": {"completeness": 85.5, "validity": 70.5, "timeliness": None},
         "rules": [{"check_id": "MM003", "name": "Material type is mandatory - x", "affected": 11_420, "total": 121_340}],
         "blockers": 1, "critical": 0}
    s = rc.module(m, T)
    assert s.startswith("Material master scores 84.1, which is below the target. It moved 1.9 points down since the earlier run.")
    assert "weakest dimension is validity at 70.5, meaning values fall outside the allowed set or pattern" in s
    assert "Material type is mandatory, failing on 11,420 records (9.4% of its records)" in s
    assert s.endswith("1 blocking rule stops this module going live as it stands.")
    assert rc.module({"name": "fi_gl", "score": None}) == "FI GL produced no score in this run."
    assert rc.module({"name": "fi_gl", "score": 95.0}, T).endswith("Nothing here blocks it; keep it there.")
    assert rc.module({"name": "fi_gl", "score": 86.0, "critical": 1, "cap_reason": "Capped at 85"}, T).endswith("1 critical failure caps its score.")
    assert rc.module({"name": "fi_gl", "score": 92.0, "critical": 1}, T).endswith("so the score is not capped, but those records cannot be used until fixed.")
    ov = rc.modules_overview([{"name": "a", "score": 95.0}, {"name": "b", "score": 70.0, "critical": 1}, {"name": "c", "score": None}], T)
    assert ov.startswith("1 of 2 scored modules meets the 90 target. B is the weakest at 70.0 and A the strongest at 95.0, a spread of 25.0 points.")
    assert "1 module carries a critical or blocking failure: B." in ov and "1 module produced no score" in ov


def test_funnel_cleaning_and_extraction():
    steps = agg.cleansing_steps(1000, 400, 300, 100, 50)
    assert rc.funnel(steps) == ("40% of failing records have a proposed value. 25% of those are approved and 50% of the approved "
                                "values are applied. The gap between approved and applied is the export batch still to be loaded.")
    assert "waiting on a steward" in rc.funnel(agg.cleansing_steps(1000, 400, 300, 0, 0))
    assert "by hand" in rc.funnel(agg.cleansing_steps(1000, 0, 0, 0, 0))
    assert rc.funnel(agg.cleansing_steps(0, 0, 0, 0, 0)).startswith("No failing records")
    assert rc.cleaning({"queue_total": 0, "recon": {}, "record_fixes_total": 0}).startswith("Nothing has been proposed")
    s = rc.cleaning({"queue_total": 65, "applied": 40, "by_status": {"pending": 25},
                     "recon": {"records": 100, "fixed": 80, "still_failing": 5, "unreconciled": 15}})
    assert s.startswith("40 of 65 proposals are applied, 62%. 25 are still pending review")
    assert "94% held and 5 are still failing" in s and "15 corrected records have not been re-checked" in s
    ctx = pr.extraction_context(fx.V2, **KW)
    s = ctx["commentary"]["run"]
    assert s.startswith("1 table failed to read (ADRC), so every check that needs it was skipped")
    assert "BSEG and MARC were read in part" in s and s.endswith("re-run before relying on it.")
    assert rc.extraction({"coverage": [{"table": "MARA", "status": "live"}], "live": 1, "failed": 0, "skipped": 0,
                          "complete": True}) == "All 1 tables were read in full, so the scores rest on complete data."


def test_trend_owner_and_object_sentences():
    d = {"name": "validity", "values": [74.0, None, 70.5], "change": -3.5}
    s = rc.trend_dimension(d)
    assert s.startswith("Validity fell 3.5 points over the window, from 74.0 to 70.5.") and s.endswith("correct the outliers.")
    assert rc.trend_dimension({"name": "accuracy", "values": [80.0], "change": None}) == ""
    assert "swung" in rc.trend_dimension({"name": "accuracy", "values": [60.0, 95.0, 70.0], "change": 10.0})
    m = rc.movers([{"module": "fi_gl", "change": 5.0}, {"module": "accounts_payable", "change": -2.0}])
    assert m == "1 object improved and 1 slipped between the first and the last run. FI GL gained most, +5.0 points. Accounts payable lost most, -2.0 points, and deserves the first look."
    assert rc.rule_changes([], []).startswith("No rule changed state")
    assert "points at a process change" in rc.rule_changes([{"check_id": "X", "module": "fi_gl", "v2": 5}], [])
    o = {"owner": "Thandi", "score": 61.0, "open": 7, "open_by_severity": {"critical": 2, "high": 5}, "trend": "slipped",
         "delta": -4.0, "oldest_item_age_days": 45, "blocked": 1, "object_count": 3}
    s = rc.owner(o)
    assert s.startswith("Thandi holds 7 open items and an owner score of 61.0. 2 are critical and should come first.")
    assert "slipped 4.0" in s and "45 days old" in s and s.endswith("1 of Thandi's 3 objects has a blocking rule.")
    assert rc.owner({"owner": "Sam", "score": None, "objects": []}).endswith("nothing to act on.")
    assert rc.unassigned([]).startswith("Every object has an owner")
    assert "Assign these first." in rc.unassigned([{"module": "fi_gl", "blocking": {"check_id": "X"}}])
    assert rc.object_rules([], []) == ""
    assert rc.samples([]) == "" and "real failing records" in rc.samples([{"check_id": "X"}])


# ── the reports carry the commentary ─────────────────────────────────────────


def test_analysis_and_executive_reports_carry_commentary_and_next_steps():
    ctx = pr.analysis_context(fx.V2, fx.FINDINGS2, system=fx.SYSTEM, previous_dqs=76.0, **KW)
    com = ctx["commentary"]
    assert set(com) == {"overall", "dimensions", "fix_first", "severity", "critical", "modules", "next_steps"}
    assert com["overall"].startswith("A score of 85.1 is below the target of 90.") and "9.1 points better" in com["overall"]
    assert com["next_steps"][0].startswith("Fix Material type is mandatory in Material master first: 11,420 records")
    paras = _paras("analysis_report.html", ctx)
    assert len(paras) == 6 and paras[0] == com["overall"] and com["fix_first"] in paras
    text = _text("analysis_report.html", ctx)
    assert "What to do next" in text and com["next_steps"][-1] in text
    supplementary = {"cleaning": {}, "dedup": {}, "exceptions": {}, "contracts": {}, "impact": {}, "mdm_snapshot": None,
                     "golden_summary": [], "dqs_trend": []}
    ex = pr.executive_context({}, supplementary, fx.V2, fx.FINDINGS2, system=fx.SYSTEM, **KW)
    assert "What to do next" in _text("executive_report.html", ex)


def test_object_comparison_and_cleaning_reports_carry_commentary():
    obj = pr.object_context("material_master", fx.V2["dqs_summary"]["material_master"],
                            [f for f in fx.FINDINGS2 if f["module"] == "material_master"], [], system=fx.SYSTEM, **KW)
    paras = _paras("object_report.html", obj)
    assert paras[0].startswith("A score of 84.1 is below the target of 90.")
    assert any(p.startswith("2 of 2 rules found failing records. The worst is Material type is mandatory") for p in paras)
    cmp_ = pr.comparison_context(fx.V1, fx.V2, fx.FINDINGS1, fx.FINDINGS2, record_diff=fx.RECORD_DIFF, **KW)
    changes = cmp_["commentary"]["changes"]
    assert changes.startswith("Completeness improved most, +8.9 points.") and "standing backlog" in changes
    assert changes in _paras("comparison_report.html", cmp_)
    cl = pr.cleaning_context({"queue": [{"object_type": "material_master", "status": "applied", "n": 40}]}, **KW)
    assert _paras("cleaning_report.html", cl) == ["40 of 40 proposals are applied, 100%."]
    assert _paras("extraction_report.html", pr.extraction_context(fx.V2, **KW))[0].startswith("1 table failed to read")


def test_board_report_commentary_every_section_and_each_module():
    ctx = pr.board_context(board_data(), **KW)
    com = ctx["commentary"]
    assert set(com) == {"overall", "dimensions", "clusters", "blockers", "waves", "funnel", "modules", "next_steps"}
    assert com["overall"].startswith("A score of") and "earlier run" in com["overall"]
    assert com["clusters"].startswith("Material master rules with no named field is the largest root cause, 71% of the impact")
    assert "30% of the top cluster's records already have a proposed value" in com["clusters"]
    assert "Wave 2 is held back" in com["waves"] or "Wave 1 is held back" in com["waves"]
    assert com["funnel"].startswith("21% of failing records have a proposed value. 22% of those are approved")
    assert com["modules"].startswith("1 of 4 scored modules meets the 90 target.")
    assert all(m["commentary"].startswith(pr.fmt_module(m["name"]) + " scores") for g in ctx["groups"] for m in g["modules"])
    text = _text("board_report.html", ctx)
    for key in ("overall", "dimensions", "clusters", "blockers", "waves", "funnel", "modules"):
        assert com[key] in text, key
    assert "What to do next" in text and com["next_steps"][0] in text
    # the module overview is printed once, under the first group only
    assert text.count(com["modules"]) == 1
    # with no earlier run nothing claims a movement
    solo = pr.board_context(board_data(with_previous=False), **KW)["commentary"]
    assert "earlier run" not in solo["overall"] and "moved most" not in solo["dimensions"]


def test_trend_and_owner_reports_carry_commentary():
    runs = agg.trend_runs([fx.V1, fx.V2], {fx.V1["id"]: fx.FINDINGS1, fx.V2["id"]: fx.FINDINGS2}, {}, {})
    ctx = pr.trend_context(fx.V2, {"runs": runs, "annotations": []}, fx.FINDINGS1, fx.FINDINGS2, system=fx.SYSTEM, **KW)
    assert any(d["commentary"] for d in ctx["dims"])
    assert ctx["commentary"]["movers"].startswith("2 objects improved and 1 slipped between the first and the last run.")
    assert ctx["commentary"]["movers"].endswith("Material master lost most, -1.8 points, and deserves the first look.")
    paras = _paras("trend_report.html", ctx)
    assert ctx["commentary"]["movers"] in paras and ctx["commentary"]["rules"] in paras
    owners = [{"user_id": "u1", "owner": "Thandi", "score": 61.0, "delta": -4.0, "open_by_severity": {"critical": 2},
               "fixed_since_baseline": 1, "oldest_item_age_days": 45}]
    objects = [{"module": "fi_gl", "score": 93.4, "owner_id": "u1", "blocking": None},
               {"module": "material_master", "score": 84.1, "owner_id": None,
                "blocking": {"check_id": "MM003", "severity": "high", "affected": 11_420, "name": "x"}}]
    ctx = pr.owners_context(fx.V2, owners, objects, system=fx.SYSTEM, **KW)
    assert ctx["owners"][0]["commentary"].startswith("Thandi holds 2 open items")
    assert ctx["commentary"]["unassigned"].startswith("1 object has no owner. 1 of them carries a blocking rule (Material master)")
    paras = _paras("owner_report.html", ctx)
    assert ctx["owners"][0]["commentary"] in paras and ctx["commentary"]["unassigned"] in paras


def test_workbook_commentary_follows_the_module_filter():
    from api.services import workbook_report as wb

    data = board_data()
    data.update(module="fi_gl", findings=[f for f in fx.FINDINGS2 if f["module"] == "fi_gl"],
                previous_findings=[f for f in fx.FINDINGS1 if f["module"] == "fi_gl"])
    data["clusters"] = agg.cluster_findings(data["findings"], {})
    data["overview"] = agg.overview_from(data)
    rows = list(wb.workbook_context(data, scoring={}, generated_at=fx.GENERATED)["sheets"][1].rows)
    text = " ".join(r["commentary"] for r in rows)
    assert "FI GL" in text and "Material master" not in text and "Wave 2" not in text
    assert [r["section"] for r in rows if r["section"] == "FI GL"] == ["FI GL"]


def test_workbook_has_a_commentary_sheet_after_the_summary():
    pytest.importorskip("openpyxl")
    from openpyxl import load_workbook

    from api.services import workbook_report as wb
    from api.services.branded_xlsx import build_workbook

    ctx = wb.workbook_context(board_data(), scoring={}, generated_at=fx.GENERATED)
    titles = [s.title for s in ctx["sheets"]]
    assert titles[:3] == ["Summary", "Commentary", "Modules"] and len(titles) == 10
    rows = list(ctx["sheets"][1].rows)
    sections = [r["section"] for r in rows]
    assert sections[:2] == ["Score", "Dimensions"] and "Material master" in sections and sections[-1].startswith("Next step")
    assert all(r["commentary"] for r in rows)
    data = build_workbook(tenant_name=fx.TENANT, run_label="r", run_id="id", title=ctx["title"],
                          sheets=ctx["sheets"], generated_at=ctx["generated_at"])
    book = load_workbook(io.BytesIO(data))
    assert "Commentary" in book.sheetnames
