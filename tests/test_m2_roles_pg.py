"""Step 6c, the role capability audit, against a real Postgres (migrations 104 and 105; docs/roles-and-features.md).

Every verb the audit added: comments and @mentions, watches and subscriptions (and the run events they fire), a
person's activity, out-of-office delegations (and the notification copy to the delegate), the escalation chain and its
worker step, contract versions, wave gates and the certificate stub, workload balancing with undo, share links,
the configuration bundle, view as a role, bulk rules, the grid export, shared views and remembered columns, the
extraction dry run, the steward assignment template, related objects and the fix-order hint. RBAC gates, audit rows
and row-level security of every new table.

Runs with MERIDIAN_TEST_DB_URL (a throwaway database migrated to head); skipped otherwise.
"""
from __future__ import annotations

import base64
import io
import os
import uuid
from datetime import date, timedelta

import pytest
from sqlalchemy import create_engine, text
from sqlalchemy.orm import Session

from tests.m2_client import headers, make_client
from tests.m2_seed import R1, seed_m2

pytestmark = pytest.mark.skipif(not os.environ.get("MERIDIAN_TEST_DB_URL"), reason="MERIDIAN_TEST_DB_URL not set")

NEW_TABLES = ("comments", "watches", "notification_subscriptions", "delegations", "escalation_steps",
              "contract_versions", "wave_gates", "share_links", "assignment_templates", "steward_assignments")


@pytest.fixture(scope="module")
def owner():
    eng = create_engine(os.environ["MERIDIAN_TEST_DB_URL"])
    yield eng
    eng.dispose()


async def client_for(monkeypatch, tid: str, sub: str):
    deps, client = make_client(monkeypatch, tid, sub)
    await deps.engine.dispose()
    return client


def _q(owner, sql: str, **p):
    with owner.begin() as c:
        return c.execute(text(sql), p).fetchall()


def _audit(owner, tid: str, action: str) -> int:
    return _q(owner, "SELECT count(*) FROM audit_log WHERE tenant_id = :t AND action = :a", t=tid, a=action)[0][0]


# ── comments, mentions, watches, subscriptions ────────────────────────────────


@pytest.mark.anyio
async def test_comment_mentions_and_watchers_are_notified(owner, monkeypatch):
    s = seed_m2(owner)
    tid, ann, bob, cara = s["tid"], s["user"], s["bob"], s["cara"]
    rec = f"material_master:{R1}"
    c = await client_for(monkeypatch, tid, cara)
    async with c:
        assert (await c.put("/api/v1/watches", json={"target_kind": "record", "target_id": rec, "events": ["comment"]},
                            headers=headers("approver"))).status_code == 200
        w = (await c.get("/api/v1/watches", params={"target_kind": "record", "target_id": rec}, headers=headers("approver"))).json()
        assert w == {"target_kind": "record", "target_id": rec, "watching": True, "events": ["comment"]}
    c = await client_for(monkeypatch, tid, ann)
    async with c:
        r = await c.post("/api/v1/comments", json={"target_kind": "record", "target_id": rec,
                                                    "body": f"Check the unit with @bob-{tid[:8]} before Friday",
                                                    "mentions": []},
                         headers=headers("steward"))
        assert r.status_code == 201, r.text
        body = r.json()
        assert body["mentions"] == [bob] and body["mine"] is True and body["author_label"]
        got = (await c.get("/api/v1/comments", params={"target_kind": "record", "target_id": rec}, headers=headers("viewer"))).json()
        assert [x["body"] for x in got["comments"]] == [f"Check the unit with @bob-{tid[:8]} before Friday"]
        assert (await c.post("/api/v1/comments", json={"target_kind": "planet", "target_id": "x", "body": "hi"},
                             headers=headers())).status_code == 422
    notes = _q(owner, "SELECT user_id::text, type, link FROM notifications WHERE tenant_id = :t AND type IN ('mention','comment') "
                      "ORDER BY type", t=tid)
    assert [(u, k) for u, k, _ in notes] == [(cara, "comment"), (bob, "mention")]
    assert all(link == f"/v2?open=record:{rec}" for _, _, link in notes)
    assert _audit(owner, tid, "comment.create") == 1
    c = await client_for(monkeypatch, tid, bob)
    async with c:
        assert (await c.get("/api/v1/comments/mentions/unread", headers=headers("steward"))).json() == {"count": 1}
        # someone else cannot delete it; the author can
        assert (await c.delete(f"/api/v1/comments/{body['id']}", headers=headers("steward"))).status_code == 403
    c = await client_for(monkeypatch, tid, ann)
    async with c:
        assert (await c.delete(f"/api/v1/comments/{body['id']}", headers=headers("steward"))).status_code == 204
        got = (await c.get("/api/v1/comments", params={"target_kind": "record", "target_id": rec}, headers=headers())).json()
        assert got["comments"] == []


@pytest.mark.anyio
async def test_subscriptions_turn_an_event_off(owner, monkeypatch):
    s = seed_m2(owner)
    tid, ann, cara = s["tid"], s["user"], s["cara"]
    c = await client_for(monkeypatch, tid, cara)
    async with c:
        subs = (await c.get("/api/v1/me/subscriptions", headers=headers("approver"))).json()["events"]
        assert {e["event"] for e in subs} >= {"mention", "comment", "run_finished", "score_dropped", "critical_breach"}
        assert all(e["enabled"] for e in subs)
        r = await c.put("/api/v1/me/subscriptions", json={"comment": False}, headers=headers("approver"))
        assert [e for e in r.json()["events"] if e["event"] == "comment"][0]["enabled"] is False
        assert (await c.put("/api/v1/me/subscriptions", json={"nonsense": True}, headers=headers())).status_code == 422
        await c.put("/api/v1/watches", json={"target_kind": "rule", "target_id": "MM014", "events": []}, headers=headers("approver"))
    c = await client_for(monkeypatch, tid, ann)
    async with c:
        await c.post("/api/v1/comments", json={"target_kind": "rule", "target_id": "MM014", "body": "Noisy rule"},
                     headers=headers())
    assert _q(owner, "SELECT count(*) FROM notifications WHERE tenant_id = :t AND type = 'comment'", t=tid)[0][0] == 0


def test_run_events_reach_object_watchers(owner):
    from api.services.collab import fire_run_events

    s = seed_m2(owner)
    tid, cara = s["tid"], s["cara"]
    with owner.begin() as c:
        c.execute(text("INSERT INTO watches (tenant_id, user_id, target_kind, target_id, events) VALUES "
                       "(:t, :u, 'object', :o, '{}')"), {"t": tid, "u": cara, "o": f"{s['ecc']}:material_master"})
    with Session(owner) as session:
        session.execute(text("SELECT set_config('app.tenant_id', :t, false)"), {"t": tid})
        sent = fire_run_events(session, tid, s["run"])
        session.commit()
    assert sent >= 1
    rows = _q(owner, "SELECT type, title, link FROM notifications WHERE tenant_id = :t AND user_id = :u ORDER BY type",
              t=tid, u=cara)
    kinds = [r[0] for r in rows]
    assert "run_finished" in kinds
    assert all(r[2] == f"/v2/o/{s['ecc']}/material_master" for r in rows)


@pytest.mark.anyio
async def test_person_activity_is_own_or_needs_audit(owner, monkeypatch):
    s = seed_m2(owner)
    c = await client_for(monkeypatch, s["tid"], s["dan"])
    async with c:
        assert (await c.get(f"/api/v1/people/{s['dan']}/activity", headers=headers("viewer"))).status_code == 200
        assert (await c.get(f"/api/v1/people/{s['bob']}/activity", headers=headers("viewer"))).status_code == 403
        r = await c.get(f"/api/v1/people/{s['bob']}/activity", headers=headers("auditor"))
        assert r.status_code == 200 and [e["action"] for e in r.json()["entries"]] == ["update"]


# ── delegations ───────────────────────────────────────────────────────────────


@pytest.mark.anyio
async def test_delegation_routes_notifications_to_the_delegate(owner, monkeypatch):
    from api.services.notifications import create_notification_sync

    s = seed_m2(owner)
    tid, bob, cara, dan = s["tid"], s["bob"], s["cara"], s["dan"]
    today = date.today()
    c = await client_for(monkeypatch, tid, cara)
    async with c:
        bad = await c.post("/api/v1/delegations", json={"to_user_id": dan, "starts_on": str(today),
                                                         "ends_on": str(today + timedelta(days=3))}, headers=headers("approver"))
        assert bad.status_code == 422 and "cannot approve" in bad.json()["detail"]
        ok = await c.post("/api/v1/delegations", json={"to_user_id": bob, "starts_on": str(today),
                                                        "ends_on": str(today + timedelta(days=3)), "reason": "Leave"},
                          headers=headers("approver"))
        assert ok.status_code == 201, ok.text
        assert ok.json()["state"] == "active" and ok.json()["to_name"] == "Bob Steward"
        again = await c.post("/api/v1/delegations", json={"to_user_id": bob, "starts_on": str(today + timedelta(days=1)),
                                                           "ends_on": str(today + timedelta(days=5))}, headers=headers("approver"))
        assert again.status_code == 409
        # someone else's out of office needs manage_users
        other = await c.post("/api/v1/delegations", json={"from_user_id": bob, "to_user_id": cara, "starts_on": str(today),
                                                           "ends_on": str(today)}, headers=headers("approver"))
        assert other.status_code == 403
        mine = (await c.get("/api/v1/delegations/mine", headers=headers("approver"))).json()
        assert len(mine["away"]) == 1 and mine["covering"] == []
        assert (await c.get("/api/v1/delegations", headers=headers("approver"))).status_code == 403
    with Session(owner) as session:
        session.execute(text("SELECT set_config('app.tenant_id', :t, false)"), {"t": tid})
        create_notification_sync(tid, cara, "approval", "Batch B-1 waits", "Approve it", "/v2/work", session)
        create_notification_sync(tid, cara, "digest", "Weekly digest", "x", None, session)
        session.commit()
    rows = _q(owner, "SELECT user_id::text, type, title FROM notifications WHERE tenant_id = :t ORDER BY title", t=tid)
    assert (bob, "approval", "For Cara Owner: Batch B-1 waits") in rows
    assert not any(u == bob and k == "digest" for u, k, _ in rows)
    c = await client_for(monkeypatch, tid, bob)
    async with c:
        assert len((await c.get("/api/v1/delegations/mine", headers=headers("steward"))).json()["covering"]) == 1
        did = ok.json()["id"]
        assert (await c.delete(f"/api/v1/delegations/{did}", headers=headers("steward"))).status_code == 204
    assert _audit(owner, tid, "delegation.create") == 1 and _audit(owner, tid, "delegation.revoke") == 1


# ── escalation chain ──────────────────────────────────────────────────────────


@pytest.mark.anyio
async def test_escalation_chain_edit_and_worker_step(owner, monkeypatch):
    from api.services.escalation import escalate_chain

    s = seed_m2(owner)
    tid, ann, cara = s["tid"], s["user"], s["cara"]
    c = await client_for(monkeypatch, tid, ann)
    async with c:
        bad = await c.put("/api/v1/triage/escalation-chain", json={"steps": [
            {"level": 1, "after_minutes": 60, "target": "fallback"}, {"level": 2, "after_minutes": 30, "target": "fallback"}]},
            headers=headers())
        assert bad.status_code == 422
        assert (await c.put("/api/v1/triage/escalation-chain", json={"steps": []}, headers=headers("viewer"))).status_code == 403
        r = await c.put("/api/v1/triage/escalation-chain", json={"steps": [
            {"level": 1, "after_minutes": 0, "target": "user", "user_id": cara}]}, headers=headers())
        assert r.status_code == 200 and r.json()["steps"][0]["user_name"] == "Cara Owner"
    with owner.begin() as conn:  # one live, breached issue for Bob
        conn.execute(text("UPDATE record_issues SET sla_started_at = now() - interval '3 days', "
                          "due_at = now() - interval '2 hours', sla_notified = '{}' WHERE tenant_id = :t AND check_id = 'MM900'"),
                     {"t": tid})
    with Session(owner) as session:
        session.execute(text("SELECT set_config('app.tenant_id', :t, false)"), {"t": tid})
        first = escalate_chain(session, tid)
        second = escalate_chain(session, tid)
        session.commit()
    assert first == {"issue:esc1": 1} and second == {}
    rows = _q(owner, "SELECT user_id::text, type FROM notifications WHERE tenant_id = :t AND type = 'escalation'", t=tid)
    assert rows == [(cara, "escalation")]
    assert _audit(owner, tid, "escalation_chain.update") == 1


# ── contracts ─────────────────────────────────────────────────────────────────


@pytest.mark.anyio
async def test_contract_revise_keeps_versions(owner, monkeypatch):
    s = seed_m2(owner)
    tid = s["tid"]
    with owner.begin() as conn:
        cid = conn.execute(text("INSERT INTO contracts (tenant_id, name, producer, consumer, status) VALUES "
                                "(:t, 'MM to S/4', 'material_master', 'S/4HANA', 'active') RETURNING id::text"), {"t": tid}).scalar()
        conn.execute(text("INSERT INTO contract_compliance_history (tenant_id, contract_id, overall_compliant, violations) "
                          "VALUES (:t, :c, false, CAST('[{\"dimension\": \"validity\"}]' AS jsonb))"), {"t": tid, "c": cid})
    c = await client_for(monkeypatch, tid, s["user"])
    async with c:
        assert (await c.post(f"/api/v1/contracts/{cid}/revise", json={}, headers=headers("approver"))).status_code == 403
        r = await c.post(f"/api/v1/contracts/{cid}/revise", json={"note": "Tighten validity"}, headers=headers("steward"))
        assert r.json() == {"id": cid, "version": 2, "status": "draft"}
        assert (await c.post(f"/api/v1/contracts/{cid}/revise", json={}, headers=headers("steward"))).status_code == 409
        d = (await c.get(f"/api/v1/contracts/{cid}/detail", headers=headers("viewer"))).json()
        assert d["contract"]["version"] == 2 and d["contract"]["status"] == "draft"
        assert [v["version"] for v in d["versions"]] == [1] and d["versions"][0]["snapshot"]["status"] == "active"
        assert len(d["breaches"]) == 1 and d["checks"] == 1


# ── wave gates ────────────────────────────────────────────────────────────────


@pytest.mark.anyio
async def test_wave_gate_status_and_certificate(owner, monkeypatch):
    s = seed_m2(owner)
    tid = s["tid"]
    with owner.begin() as conn:
        wid = conn.execute(text("INSERT INTO migration_waves (tenant_id, name, modules) VALUES (:t, 'Wave 1', "
                                "ARRAY['material_master','fi_gl']) RETURNING id::text"), {"t": tid}).scalar()
    c = await client_for(monkeypatch, tid, s["user"])
    async with c:
        g = (await c.get("/api/v1/migration/waves/gates", headers=headers("viewer"))).json()["gates"]
        assert g[0]["gate"]["configured"] is False
        assert (await c.put(f"/api/v1/migration/waves/{wid}/gate", json={"min_score": 50}, headers=headers("viewer"))).status_code == 403
        r = await c.put(f"/api/v1/migration/waves/{wid}/gate", json={"min_score": 99.5, "zero_criticals": True,
                                                                        "contracts_met": True}, headers=headers("manager"))
        assert r.status_code == 200, r.text
        st = r.json()["status"]
        assert st["passed"] is False and [x["key"] for x in st["checks"]] == ["min_score", "zero_criticals", "contracts_met"]
        assert st["checks"][2]["passed"] is True  # no active contract names these objects
        cert = (await c.get(f"/api/v1/migration/waves/{wid}/certificate", headers=headers("viewer"))).json()
        assert cert["final"] is False and cert["passed"] is False and cert["wave"] == "Wave 1"
    assert _audit(owner, tid, "wave_gate.update") == 1


# ── workload balancing ────────────────────────────────────────────────────────


@pytest.mark.anyio
async def test_balance_plan_apply_and_undo(owner, monkeypatch):
    s = seed_m2(owner)
    tid, ann, bob = s["tid"], s["user"], s["bob"]
    c = await client_for(monkeypatch, tid, ann)
    async with c:
        assert (await c.post("/api/v1/people/balance/plan", json={}, headers=headers("viewer"))).status_code == 403
        plan = await c.post("/api/v1/people/balance/plan", json={"by": "owner", "user_ids": [ann, bob]}, headers=headers("manager"))
        assert plan.status_code == 200, plan.text
        p = plan.json()
        assert sum(p["after"].values()) == sum(p["before"].values()) + sum(m["count"] for m in p["moves"] if m["from"] is None)
        assert max(p["after"].values()) - min(p["after"].values()) <= 1
        r = await c.post("/api/v1/people/balance/apply", json={"by": "owner", "user_ids": [ann, bob], "moves": p["moves"]},
                         headers=headers("manager"))
        assert r.status_code == 200 and r.json()["moved"] == p["moved"] and r.json()["action_id"]
        u = await c.post("/api/v1/actions/undo", json={"action_id": r.json()["action_id"]}, headers=headers("manager"))
        assert u.status_code == 200, u.text and u.json()["restored"]["issues"] == p["moved"]


# ── share links ───────────────────────────────────────────────────────────────


@pytest.mark.anyio
async def test_share_link_lifecycle(owner, monkeypatch):
    s = seed_m2(owner)
    tid, ann, dan = s["tid"], s["user"], s["dan"]
    c = await client_for(monkeypatch, tid, ann)
    async with c:
        assert (await c.post("/api/v1/share-links", json={"route": "/v2"}, headers=headers("viewer"))).status_code == 403
        assert (await c.post("/api/v1/share-links", json={"route": "/admin"}, headers=headers("auditor"))).status_code == 422
        r = await c.post("/api/v1/share-links", json={"route": "/v2/o/ecc-production/material-master",
                                                       "state": {"tab": "rules", "fixture": "1"}, "expires_in_days": 3},
                         headers=headers("auditor"))
        assert r.status_code == 201, r.text
        link = r.json()
        assert link["url"] == f"/v2/share/{link['token']}" and link["state"] == {"tab": "rules"}
        stored = _q(owner, "SELECT token_hash FROM share_links WHERE id = :i", i=link["id"])[0][0]
        assert link["token"] not in stored and len(stored) == 64
    c = await client_for(monkeypatch, tid, dan)
    async with c:
        res = await c.get(f"/api/v1/share-links/resolve/{link['token']}", headers=headers("viewer"))
        assert res.status_code == 200 and res.json()["role"] == "viewer" and res.json()["route"].endswith("material-master")
        assert (await c.get("/api/v1/share-links", params={"all": True}, headers=headers("viewer"))).status_code == 403
        assert (await c.delete(f"/api/v1/share-links/{link['id']}", headers=headers("viewer"))).status_code == 403
    c = await client_for(monkeypatch, tid, ann)
    async with c:
        assert len((await c.get("/api/v1/share-links", params={"all": True}, headers=headers("admin"))).json()["links"]) == 1
        assert (await c.delete(f"/api/v1/share-links/{link['id']}", headers=headers("admin"))).status_code == 204
        assert (await c.get(f"/api/v1/share-links/resolve/{link['token']}", headers=headers())).status_code == 410
    assert _q(owner, "SELECT open_count FROM share_links WHERE id = :i", i=link["id"])[0][0] == 1
    assert _audit(owner, tid, "share_link.open") == 1 and _audit(owner, tid, "share_link.revoke") == 1


# ── configuration bundle, view as, bulk rules, grid export ────────────────────


@pytest.mark.anyio
async def test_config_bundle_round_trip(owner, monkeypatch):
    s = seed_m2(owner)
    tid = s["tid"]
    c = await client_for(monkeypatch, tid, s["user"])
    async with c:
        assert (await c.get("/api/v1/admin/config-bundle", headers=headers("steward"))).status_code == 403
        b = (await c.get("/api/v1/admin/config-bundle", headers=headers())).json()
        assert b["format"] == "meridian.tenant-config" and set(b["sections"]) == {
            "settings", "rules", "mappings", "waves", "sla_policies", "escalation"}
        same = (await c.post("/api/v1/admin/config-bundle/import", json=b, headers=headers())).json()
        assert same["dry_run"] is True and same["totals"]["add"] == 0 and same["totals"]["change"] == 0
        b["sections"]["waves"].append({"name": "Wave 9", "modules": ["fi_gl"], "stage": "plan", "gate": {"min_score": 80}})
        b["sections"]["sla_policies"] = [{"severity": "critical", "module": None, "ack_minutes": 30, "resolve_minutes": 600}]
        dry = (await c.post("/api/v1/admin/config-bundle/import", json=b, headers=headers())).json()
        assert dry["sections"]["waves"]["add"] == 1 and dry["sections"]["sla_policies"]["change"] == 1
        assert dry["sections"]["sla_policies"]["kept"] == 1
        assert _q(owner, "SELECT count(*) FROM migration_waves WHERE tenant_id = :t", t=tid)[0][0] == 0
        done = await c.post("/api/v1/admin/config-bundle/import", params={"dry_run": False}, json=b, headers=headers())
        assert done.status_code == 200 and done.json()["dry_run"] is False
        assert (await c.post("/api/v1/admin/config-bundle/import", json={"format": "x"}, headers=headers())).status_code == 422
    assert _q(owner, "SELECT g.min_score FROM migration_waves w JOIN wave_gates g ON g.wave_id = w.id WHERE w.tenant_id = :t",
              t=tid) == [(80.0,)]
    assert _q(owner, "SELECT resolve_minutes FROM sla_policies WHERE tenant_id = :t AND severity = 'critical'", t=tid) == [(600,)]
    assert _audit(owner, tid, "config_bundle.import") == 1


@pytest.mark.anyio
async def test_view_as_narrows_and_is_read_only(owner, monkeypatch):
    s = seed_m2(owner)
    tid = s["tid"]
    c = await client_for(monkeypatch, tid, s["user"])
    async with c:
        r = await c.post("/api/v1/admin/view-as", json={"role": "data_steward"}, headers=headers())
        assert r.status_code == 200 and r.json()["read_only"] is True and "manage_users" not in r.json()["actions"]
        assert (await c.post("/api/v1/admin/view-as", json={"role": "viewer"}, headers=headers("steward"))).status_code == 403
        as_viewer = {**headers(), "X-View-As-Role": "viewer"}
        me = (await c.get("/api/v1/me/permissions", headers=as_viewer)).json()
        assert me["role"] == "viewer" and "export" not in me["actions"]
        assert (await c.get("/api/v1/admin/config-bundle", headers=as_viewer)).status_code == 403
        w = await c.put("/api/v1/watches", json={"target_kind": "rule", "target_id": "MM014"}, headers=as_viewer)
        assert w.status_code == 403 and "viewing as" in w.json()["detail"]
        # a non-administrator cannot use the header at all
        assert (await c.get("/api/v1/landscape", headers={**headers("steward"), "X-View-As-Role": "viewer"})).status_code == 403
        assert (await c.delete("/api/v1/admin/view-as", headers=as_viewer)).status_code == 204
    assert _audit(owner, tid, "view_as.start") == 1 and _audit(owner, tid, "view_as.stop") == 1


@pytest.mark.anyio
async def test_bulk_rules_import_validates_and_export(owner, monkeypatch):
    s = seed_m2(owner)
    c = await client_for(monkeypatch, s["tid"], s["user"])
    async with c:
        bad = "- module: material_master\n  id: ZX1\n  severity: extreme\n- [not, a, rule]\n"
        assert (await c.post("/api/v1/rules/bulk/import", json={"format": "yaml", "content": bad},
                             headers=headers("viewer"))).status_code == 403
        r = await c.post("/api/v1/rules/bulk/import", json={"format": "yaml", "content": bad}, headers=headers("steward"))
        assert r.status_code == 200, r.text
        out = r.json()
        assert out["dry_run"] is True and out["valid"] == 0 and out["invalid"] == 2 and out["created"] == []
        assert any("severity" in e for e in out["rows"][0]["errors"])
        assert (await c.post("/api/v1/rules/bulk/import", json={"format": "yaml", "content": "a: ["},
                             headers=headers("steward"))).status_code == 422
        y = await c.get("/api/v1/rules/bulk/export", params={"format": "yaml"}, headers=headers("auditor"))
        assert y.status_code == 200 and y.headers["content-type"].startswith("application/x-yaml")
        x = await c.get("/api/v1/rules/bulk/export", params={"format": "xlsx"}, headers=headers("auditor"))
        assert x.status_code == 200 and x.content[:2] == b"PK"
        back = await c.post("/api/v1/rules/bulk/import", json={"format": "xlsx", "content": base64.b64encode(x.content).decode()},
                            headers=headers("steward"))
        assert back.status_code in (200, 422)  # an empty catalogue exports a sheet with no rules


@pytest.mark.anyio
async def test_grid_export_writes_the_visible_columns(owner, monkeypatch):
    from openpyxl import load_workbook

    s = seed_m2(owner)
    c = await client_for(monkeypatch, s["tid"], s["user"])
    async with c:
        body = {"title": "People", "route": "people", "columns": [{"key": "name", "header": "Name"},
                                                                 {"key": "open", "header": "Open work", "kind": "int"}],
                "rows": [{"name": "=HYPERLINK(1)", "open": 3, "hidden": "x"}], "filters": {"role": "approver"}}
        assert (await c.post("/api/v1/export/grid", json=body, headers=headers("viewer"))).status_code == 403
        r = await c.post("/api/v1/export/grid", json=body, headers=headers("auditor"))
        assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.worksheets[1]
    values = [[c.value for c in row] for row in ws.iter_rows()]
    flat = [v for row in values for v in row if v is not None]
    assert "Name" in flat and "Open work" in flat and "x" not in flat
    assert not any(isinstance(v, str) and v.startswith("=") for v in flat)


# ── views ─────────────────────────────────────────────────────────────────────


@pytest.mark.anyio
async def test_shared_views_and_remembered_columns(owner, monkeypatch):
    s = seed_m2(owner)
    tid = s["tid"]
    c = await client_for(monkeypatch, tid, s["bob"])
    async with c:
        await c.put("/api/v1/views", json={"route": "people", "name": "Team view", "filters": {"role": "approver"},
                                          "shared": True}, headers=headers("steward"))
        await c.put("/api/v1/views", json={"route": "people", "name": "_columns", "columns": ["name", "open"]},
                    headers=headers("steward"))
        assert (await c.put("/api/v1/views", json={"route": "people", "name": "_columns", "shared": True},
                            headers=headers("steward"))).status_code == 422
        mine = (await c.get("/api/v1/views", params={"route": "people"}, headers=headers("steward"))).json()
        assert mine["columns"] == ["name", "open"] and [v["name"] for v in mine["views"] if not v["builtin"]] == ["Team view"]
    c = await client_for(monkeypatch, tid, s["dan"])
    async with c:
        theirs = (await c.get("/api/v1/views", params={"route": "people"}, headers=headers("viewer"))).json()
        shared = [v for v in theirs["views"] if not v["builtin"]]
        assert [(v["name"], v["shared"], v["mine"]) for v in shared] == [("Team view", True, False)]
        assert theirs["columns"] is None
        assert (await c.delete(f"/api/v1/views/{shared[0]['id']}", headers=headers("viewer"))).status_code == 404


# ── extraction dry run ────────────────────────────────────────────────────────


@pytest.mark.anyio
async def test_extraction_dry_run_reads_nothing(owner, monkeypatch):
    s = seed_m2(owner)
    tid, ecc = s["tid"], s["ecc"]
    with owner.begin() as conn:
        conn.execute(text("INSERT INTO table_profiles (tenant_id, version_id, system_id, table_name, row_count) "
                          "VALUES (:t, :v, :s, 'MARA', 61816)"), {"t": tid, "v": s["run"], "s": ecc})
    queued = []
    monkeypatch.setattr("workers.tasks.run_extraction.run_extraction.delay", lambda *a, **k: queued.append((a, k)))
    c = await client_for(monkeypatch, tid, s["user"])
    async with c:
        r = await c.post(f"/api/v1/systems/{ecc}/downloads", json={"objects": ["material_master"], "dry_run": True},
                         headers=headers("sap_basis"))
        assert r.status_code in (200, 202), r.text
        out = r.json()
        assert out["dry_run"] is True and queued == []
        mara = [t for t in out["tables"] if t["table"] == "MARA"]
        assert mara and mara[0]["rows"] == 61816 and out["total_rows"] >= 61816
        assert out["rate_basis"] == "assumed" and out["seconds"] >= 1
        assert out["delta"]["available"] is False and out["delta"]["reason"]
        assert (await c.post(f"/api/v1/systems/{ecc}/downloads", json={"objects": ["material_master"], "dry_run": True},
                             headers=headers("viewer"))).status_code == 403
        live = await c.post(f"/api/v1/systems/{ecc}/downloads", json={"objects": ["material_master"], "delta": True},
                            headers=headers("sap_basis"))
        assert live.status_code == 202 and queued and queued[0][1] == {"delta": True}


# ── steward assignment template and related objects ───────────────────────────


@pytest.mark.anyio
async def test_related_objects_and_fix_hint(owner, monkeypatch):
    s = seed_m2(owner)
    c = await client_for(monkeypatch, s["tid"], s["user"])
    async with c:
        r = (await c.get("/api/v1/objects/mm_purchasing/related", headers=headers("viewer"))).json()
        rel = {x["module"]: x["relation"] for x in r["related"]}
        assert rel["material_master"] == "upstream" and rel["accounts_payable"] == "upstream"
        assert r["fix_hint"]["module"] == "material_master" and r["fix_hint"]["sentence"].startswith("Fix Material master first: ")
        mm = (await c.get("/api/v1/objects/material_master/related", headers=headers("viewer"))).json()
        assert {x["module"]: x["relation"] for x in mm["related"]}["mm_purchasing"] == "downstream"


@pytest.mark.anyio
async def test_assignment_template_apply_and_gaps(owner, monkeypatch):
    s = seed_m2(owner)
    tid, ann, bob, cara, dan = s["tid"], s["user"], s["bob"], s["cara"], s["dan"]
    c = await client_for(monkeypatch, tid, ann)
    async with c:
        t = (await c.get("/api/v1/admin/assignment-template", headers=headers("viewer"))).json()
        assert t["saved"] is False and t["template"]["objects"]["material_master"]["slice"] == "plant"
        assert t["active_slices"]["plant"] == ["3000"]
        gaps = (await c.get("/api/v1/admin/assignment-template/gaps", headers=headers("viewer"))).json()["gaps"]
        assert any(g["module"] == "fi_gl" and g["gap"] == "no_owner" for g in gaps)
        tmpl = t["template"]
        tmpl["objects"]["material_master"].update(owner_user_id=cara, steward_user_id=bob)
        tmpl["objects"]["fi_gl"].update(owner_user_id=dan, steward_user_id=bob)
        assert (await c.put("/api/v1/admin/assignment-template", json={"template": tmpl}, headers=headers("steward"))).status_code == 403
        tmpl_bad = {"objects": {"fi_gl": {"owner_role": "", "slice": "galaxy"}}}
        assert (await c.put("/api/v1/admin/assignment-template", json={"template": tmpl_bad}, headers=headers())).status_code == 422
        assert (await c.put("/api/v1/admin/assignment-template", json={"template": tmpl}, headers=headers())).status_code == 200
        dry = (await c.post("/api/v1/admin/assignment-template/apply", json={}, headers=headers())).json()
        assert dry["dry_run"] is True
        own = {r["module"]: r["change"] for r in dry["owners"]}
        assert own == {"material_master": "same", "fi_gl": "add"}
        st = {(r["module"], r["slice_value"]): r["change"] for r in dry["stewards"]}
        assert st == {("material_master", "3000"): "add", ("fi_gl", "*"): "add"}
        assert _q(owner, "SELECT count(*) FROM steward_assignments WHERE tenant_id = :t", t=tid)[0][0] == 0
        done = await c.post("/api/v1/admin/assignment-template/apply", json={"dry_run": False}, headers=headers())
        assert done.status_code == 200
        gaps = (await c.get("/api/v1/admin/assignment-template/gaps", headers=headers("viewer"))).json()["gaps"]
        assert not [g for g in gaps if g["module"] in ("material_master", "fi_gl")]
        plan = (await c.post("/api/v1/people/balance/plan", json={"by": "slice", "user_ids": [ann, bob]},
                             headers=headers("manager"))).json()
        assert plan["by"] == "slice" and all(m["to"] == bob for m in plan["moves"])
    # an owner applies only to the objects they own
    c = await client_for(monkeypatch, tid, cara)
    async with c:
        assert (await c.post("/api/v1/admin/assignment-template/apply", json={"modules": ["material_master"]},
                             headers=headers("approver"))).status_code == 200
        assert (await c.post("/api/v1/admin/assignment-template/apply", json={"modules": ["fi_gl"]},
                             headers=headers("approver"))).status_code == 403
    c = await client_for(monkeypatch, tid, dan)
    async with c:
        assert (await c.post("/api/v1/admin/assignment-template/apply", json={}, headers=headers("viewer"))).status_code == 403
    assert _audit(owner, tid, "assignment_template.apply") == 1


# ── row-level security ────────────────────────────────────────────────────────


def test_rls_isolates_every_new_table(owner):
    a, b = str(uuid.uuid4()), str(uuid.uuid4())
    with owner.begin() as c:
        for t in (a, b):
            c.execute(text("INSERT INTO tenants (id, name) VALUES (:t, 'R')"), {"t": t})
            u = str(uuid.uuid4())
            c.execute(text("INSERT INTO users (id, tenant_id, email, name, role) VALUES (:i, :t, :e, 'u', 'viewer')"),
                      {"i": u, "t": t, "e": f"r-{u[:8]}@example.test"})
            u2 = str(uuid.uuid4())
            c.execute(text("INSERT INTO users (id, tenant_id, email, name, role) VALUES (:i, :t, :e, 'v', 'viewer')"),
                      {"i": u2, "t": t, "e": f"r-{u2[:8]}@example.test"})
            cid = c.execute(text("INSERT INTO contracts (tenant_id, name, producer, consumer) VALUES (:t, 'c', 'p', 'q') "
                                 "RETURNING id"), {"t": t}).scalar()
            wid = c.execute(text("INSERT INTO migration_waves (tenant_id, name) VALUES (:t, 'w') RETURNING id"), {"t": t}).scalar()
            p = {"t": t, "u": u, "u2": u2, "c": cid, "w": wid}
            c.execute(text("INSERT INTO comments (tenant_id, target_kind, target_id, body, author_id, author_label) "
                           "VALUES (:t, 'rule', 'MM014', 'x', :u, 'u')"), p)
            c.execute(text("INSERT INTO watches (tenant_id, user_id, target_kind, target_id) VALUES (:t, :u, 'rule', 'MM014')"), p)
            c.execute(text("INSERT INTO notification_subscriptions (tenant_id, user_id, event, enabled) VALUES (:t, :u, 'comment', false)"), p)
            c.execute(text("INSERT INTO delegations (tenant_id, from_user_id, to_user_id, starts_on, ends_on, created_by) "
                           "VALUES (:t, :u, :u2, current_date, current_date, 'x')"), p)
            c.execute(text("INSERT INTO escalation_steps (tenant_id, level, after_minutes, target) VALUES (:t, 1, 0, 'fallback')"), p)
            c.execute(text("INSERT INTO contract_versions (tenant_id, contract_id, version, snapshot, created_by) "
                           "VALUES (:t, :c, 1, '{}', 'x')"), p)
            c.execute(text("INSERT INTO wave_gates (wave_id, tenant_id) VALUES (:w, :t)"), p)
            c.execute(text("INSERT INTO share_links (tenant_id, token_hash, route, created_by, expires_at) "
                           "VALUES (:t, md5(random()::text), '/v2', 'x', now() + interval '1 day')"), p)
            c.execute(text("INSERT INTO assignment_templates (tenant_id, template) VALUES (:t, '{}')"), p)
            c.execute(text("INSERT INTO steward_assignments (tenant_id, module, slice_kind, slice_value, steward_user_id) "
                           "VALUES (:t, 'fi_gl', 'none', '*', :u)"), p)
        c.execute(text("DO $$ BEGIN IF NOT EXISTS (SELECT 1 FROM pg_roles WHERE rolname = 'm2_roles_rls_app') THEN "
                       "CREATE ROLE m2_roles_rls_app NOSUPERUSER NOBYPASSRLS; END IF; END $$"))
        for t in NEW_TABLES:
            c.execute(text(f"GRANT SELECT, UPDATE ON {t} TO m2_roles_rls_app"))
    with owner.connect() as conn:
        tx = conn.begin()
        conn.execute(text("SET LOCAL ROLE m2_roles_rls_app"))
        conn.execute(text("SELECT set_config('app.tenant_id', :t, true)"), {"t": a})
        for t in NEW_TABLES:
            tenants = {str(r[0]) for r in conn.execute(text(f"SELECT tenant_id FROM {t}"))}
            assert tenants == {a}, t
            assert conn.execute(text(f"UPDATE {t} SET tenant_id = tenant_id WHERE tenant_id = :b"), {"b": b}).rowcount == 0, t
        tx.rollback()
