"""SAP correction packages: files only, never a call to SAP."""
import io
import zipfile

from openpyxl import load_workbook

from api.services import sap_packages as pk
from sap.ddic import get_dictionary

ITEMS = [
    {"module": "material_master", "check_id": "LR-000001", "record_key": "MATNR=000000000000000042",
     "field": "MARA.MATKL", "current_value": "misc", "proposed_value": "MG-0001"},
    {"module": "material_master", "check_id": "LR-000002", "record_key": "MATNR=000000000000000042|WERKS=1000",
     "field": "MARC.EKGRP", "current_value": "", "proposed_value": "=cmd|' /C calc'!A0"},
    {"module": "accounts_payable", "check_id": "AP_T", "record_key": "LIFNR=0000100002",
     "field": "LFA1.LAND1", "current_value": None, "proposed_value": "DE"},
    {"module": "x", "check_id": "N", "record_key": "K=1", "field": "ZTAB.F", "current_value": None,
     "proposed_value": None},
]


def test_ltmc_workbook_groups_tables_into_object_sheets_and_blocks_formulas():
    wb = load_workbook(io.BytesIO(pk.ltmc_workbook(ITEMS, get_dictionary("s4hana"))))
    material_sheets = [s for s in wb.sheetnames if s.startswith("Product")]
    assert material_sheets == ["Product - Basic Data", "Product - Plant Data"]
    plant = wb["Product - Plant Data"]
    header = [c.value for c in plant[1]]
    assert header[:2] == ["MATNR", "WERKS"] and "EKGRP" in header
    assert all(c.data_type != "f" for row in plant.iter_rows() for c in row)
    assert "Supplier - General Data" in wb.sheetnames
    assert "README" in wb.sheetnames                       # names the object, release check, "file only"


def test_mass_maintenance_zip_groups_by_field_and_value():
    z = zipfile.ZipFile(io.BytesIO(pk.mass_maintenance_zip(ITEMS)))
    names = sorted(z.namelist())
    assert names == ["MM17/MARA-MATKL-001.txt", "MM17/MARC-EKGRP-001.txt", "README.txt", "XK99/LFA1-LAND1-001.txt",
                     "changes.tsv"]
    assert z.read("MM17/MARA-MATKL-001.txt").decode() == "000000000000000042\n"
    changes = z.read("changes.tsv").decode().splitlines()
    assert changes[0] == "TCODE\tTABLE\tFIELD\tNEW_VALUE\tRECORD_KEY\tOLD_VALUE\tRULE"
    assert "MG-0001" in z.read("README.txt").decode()
    assert pk.mass_maintenance_zip(ITEMS) == pk.mass_maintenance_zip(list(reversed(ITEMS)))   # deterministic
